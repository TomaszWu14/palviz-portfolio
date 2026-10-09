"""Generator SYNTETYCZNEJ hali „B0” — dane w 100% fikcyjne.

Hala B0 w tym repo to wymyślony magazyn: liczba przejść, gniazd, profile regałów,
wysokości i udźwigi są zmyślone. Format plików (kolumny, arkusze, kody lokalizacji
w stylu `B0-01-100A`, `B0-07-300C-1`) odpowiada temu, czego oczekują importery i testy.

Generuje (deterministycznie, bez losowości):
  web/ui/data/warehouse_layout_b0.xlsx    — rysunek elewacji (seed „Układ B0 (wzorcowy)”)
  web/ui/data/warehouse_map_b0.xlsx       — plan fizyczny (manage.py import_locations --map)
  web/ui/data/warehouse_locations_b0.csv  — lista lokalizacji w formacie eksportu SAP
  web/wh3d/tests/data/ewm_b0_sample.csv   — próbka eksportu EWM dla testów „Wykryj z EWM”

Uruchomienie (z katalogu głównego repo):
    python web/ui/data/generate_synthetic_b0.py

Ograniczenia narzucone przez testy (nie zmieniaj bez poprawienia testów):
  * układ: dokładnie 15 przejść, > 5000 lokalizacji, poziomy {1, 2, 3, 4}, półki C/D
    (też dzielone -1/-2) — wh3d/tests/test_warehouse_layout.py;
  * próbka EWM: przejścia/gniazda z wh3d/tests/ewm_sample.py (SAMPLE_N_BAYS), 18 szablonów
    gniazd, 3327 kodów, numeracje i wyjątki z wh3d/tests/test_ewm_detect.py i test_ewm_service.py.
"""
import re
from pathlib import Path

import openpyxl

WEB = Path(__file__).resolve().parents[2]
UI_DATA = WEB / "ui" / "data"
EWM_SAMPLE = WEB / "wh3d" / "tests" / "data" / "ewm_b0_sample.csv"

# Litera → poziom fizyczny (zgodnie z ui.views.core.ewm_levels: B/C/D = półki poz. 1,
# G/H = X podzielone pionowo, S–W = strefa kompaktowa poziomy 1–5).
LEVEL = {"A": 1, "B": 1, "C": 1, "D": 1, "S": 1, "X": 2, "G": 2, "H": 2, "T": 2,
         "Y": 3, "U": 3, "Z": 4, "V": 4, "W": 5}
LETTER_ORDER = "ABCDSXGHTYUZVW"          # od podłogi w górę
CODE_RE = re.compile(r"^B0-(\d{2})-(\d{2})(\d)([A-Z])(?:-([12]))?$")

# ---------------------------------------------------------------- układ hali
# Profile regałów: lista liter (z połówkami „C-1”/„C-2”) w jednym stosie.
PALETOWY = ["A", "X", "Y", "Z"]
POLKOWY = ["B", "C", "D", "X", "Y", "Z"]
DZIELONY = ["B", "C-1", "C-2", "D-1", "D-2", "X", "Y", "Z"]
MIESZANY = ["B", "C", "G", "H", "Y", "Z"]
PRZEJAZD = ["Y", "Z"]                     # gniazda nad drogą przejazdu: tylko górne poziomy

# przejście → (ostatnie gniazdo, pozycji na belce, profil, gniazda przejazdu)
AISLES = {
    "01": (40, 3, PALETOWY, ()),
    "02": (44, 3, POLKOWY, ()),
    "03": (52, 3, PALETOWY, (25, 26)),
    "04": (42, 2, DZIELONY, ()),
    "05": (36, 3, PALETOWY, ()),
    "07": (50, 3, DZIELONY, ()),
    "08": (48, 3, PALETOWY, ()),
    "09": (38, 2, POLKOWY, ()),
    "12": (46, 3, PALETOWY, (25, 26)),
    "15": (44, 3, POLKOWY, ()),
    "28": (48, 2, DZIELONY, ()),
    "38": (52, 3, PALETOWY, (25, 26)),
    "39": (40, 3, POLKOWY, ()),
    "55": (34, 3, MIESZANY, ()),
    "76": (30, 3, PALETOWY, ()),
}
FIRST_BAY = 10
# Strefa kompaktowa S–W (poziomy 1–5) — tylko w liście lokalizacji i planie, nie w rysunku
# elewacji (seed układu ma poziomy 1–4).
COMPACT = {"76": [(70, p) for p in range(3)]}

# Kody, na które powołują się testy i kod (fikstury) — dokładane, by istniały też w danych.
REFERENCED = """
B0-01-100A B0-01-100B B0-01-100C B0-01-100C-1 B0-01-100D B0-01-100X B0-01-100Y B0-01-100Z
B0-01-101A B0-01-101X B0-01-101Y B0-01-102A B0-01-103A B0-01-104A B0-01-105A B0-01-105C
B0-01-106A B0-01-130A B0-01-130B B0-01-131A B0-01-200A B0-01-200X B0-01-200Y B0-01-300A
B0-01-300B B0-01-300C B0-01-300C-1 B0-01-300C-2 B0-01-300D B0-01-300D-2 B0-01-300G
B0-01-300H B0-01-300X B0-01-300Y B0-01-300Z B0-01-301A B0-01-301C-1 B0-01-301C-2 B0-01-301X
B0-01-301Z B0-01-302C-2 B0-01-302X B0-01-400A B0-02-100A B0-02-100X B0-02-100Y B0-02-110A
B0-02-200B B0-02-200C B0-02-300A B0-02-300B B0-02-300C B0-02-302A B0-02-302C-1 B0-02-400D
B0-03-100A B0-03-100X B0-03-100Y B0-03-100Z B0-03-120B B0-03-300A B0-03-300C B0-03-500A
B0-04-100A B0-04-400D B0-05-100A B0-07-290A B0-07-300A B0-07-300B B0-07-300C B0-07-300C-1
B0-07-300C-2 B0-07-300D-1 B0-07-300D-2 B0-07-300X B0-07-300Y B0-07-301A B0-07-301B
B0-07-301C B0-07-301C-1 B0-07-302B B0-07-302D-2 B0-07-302Z B0-07-303A B0-07-303Y B0-07-310B
B0-07-312X B0-07-312Z B0-07-480D-1 B0-07-480D-2 B0-07-700A B0-08-221A B0-08-221X B0-08-221Y
B0-08-221Z B0-09-100A B0-12-100A B0-15-422B B0-28-471D B0-38-471A B0-38-741A B0-39-300B
B0-39-300C B0-39-300D B0-39-300X B0-39-300Y B0-39-300Z B0-55-300A B0-55-700B B0-55-700C
B0-55-700G B0-55-700H B0-55-700Y B0-55-700Z B0-76-700S B0-76-700T B0-76-700U B0-76-700V
B0-76-700W
""".split()


def hall_codes():
    """Wszystkie kody syntetycznej hali (bez duplikatów, posortowane)."""
    codes = set(REFERENCED)
    for aisle, (last, k, profile, passage) in AISLES.items():
        for bay in range(FIRST_BAY, last + 1):
            letters = PRZEJAZD if bay in passage else profile
            codes |= {f"B0-{aisle}-{bay:02d}{p}{L}" for p in range(k) for L in letters}
    for aisle, stacks in COMPACT.items():
        codes |= {f"B0-{aisle}-{bay:02d}{p}{L}" for bay, p in stacks for L in "STUVW"}
    return sorted(codes, key=sort_key)


def sort_key(code):
    m = CODE_RE.match(code)
    return m.group(1), int(m.group(2)), int(m.group(3)), LETTER_ORDER.index(m.group(4)), m.group(5) or ""


def parts(code):
    """Kod → (przejście, stos „bay+pozycja”, litera, połówka|None)."""
    m = CODE_RE.match(code)
    return m.group(1), m.group(2) + m.group(3), m.group(4), m.group(5)


def location_attrs(code, i):
    """Fikcyjne atrybuty SAP: (typ magazynu, wysokość cm, objętość m³, waga kg)."""
    _a, _s, letter, half = parts(code)
    if letter in "STUVW":
        return "0012", 40, 0.35, 100
    if letter == "A":
        return ("0050" if i % 17 == 0 else "0052"), 150, 1.44, 1000
    if letter in "BCD":
        h = 120 if letter == "B" else 70
        width = 0.4 if half else 0.8
        return "0011", h, round(width * 1.2 * h / 100, 2), (150 if half else 300)
    if letter in "GH":
        return "0010", 80, 0.77, 450
    return "0010", 170, 1.63, 900


def stack_columns(codes):
    """Stos → kolumna arkusza: wspólna oś wzdłuż przejść (ten sam numer stosu = ta sama kolumna)."""
    stacks = sorted({parts(c)[1] for c in codes})
    return {s: 2 + i for i, s in enumerate(stacks)}


# ------------------------------------------------------------------ zapis plików
def write_locations_csv(codes, path):
    lines = ["Adres lokalizacji;Poziom;Typ magazynu;Wysokość [cm];Max objętość [m³];Max waga [kg]"]
    for i, code in enumerate(codes):
        typ, h, vol, kg = location_attrs(code, i)
        lines.append(f"{code};{LEVEL[parts(code)[2]]};{typ};{h};{str(vol).replace('.', ',')};{kg}")
    path.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8-sig")


def write_map_xlsx(codes, path):
    """Plan fizyczny: jeden kod w komórce; stos = kolumna, poziomy w górę od wiersza podłogi."""
    rows_per_letter = ["A", "B", "C", "C-1", "C-2", "D", "D-1", "D-2", "S", "X", "G", "H", "T",
                       "Y", "U", "Z", "V", "W"]
    band = len(rows_per_letter) + 1
    aisles = sorted({parts(c)[0] for c in codes})
    cols = stack_columns(codes)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MAPA B0 (syntetyczna)"
    ws.cell(1, 1, "Hala B0 — plan syntetyczny (dane fikcyjne)")
    ws.cell(2, 1, "LOKALIZACJE WYŁĄCZONE (GAŚNICE)")
    for ai, aisle in enumerate(aisles):
        floor = 4 + ai * band + band - 1
        ws.cell(floor, 1, f"Przejście {aisle}")
        for code in (c for c in codes if parts(c)[0] == aisle):
            _a, stack, letter, half = parts(code)
            key = letter + (f"-{half}" if half else "")
            ws.cell(floor - rows_per_letter.index(key), cols[stack], code)
    wb.save(path)


def write_layout_xlsx(codes, path):
    """Rysunek elewacji: stos = kolumna, półki w kolejnych wierszach; połówki -1/-2 w jednej
    komórce („… | …”). Strefa kompaktowa S–W pominięta (seed ma poziomy 1–4)."""
    codes = [c for c in codes if parts(c)[2] not in "STUVW"]
    letter_rows = ["Z", "Y", "H", "G", "X", "D", "C", "B", "A"]          # od góry rysunku
    band = len(letter_rows) + 2
    aisles = sorted({parts(c)[0] for c in codes})
    cols = stack_columns(codes)
    cells = {}
    for code in codes:
        aisle, stack, letter, _half = parts(code)
        r = 3 + aisles.index(aisle) * band + 1 + letter_rows.index(letter)
        cells.setdefault((r, cols[stack]), []).append(code)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Arkusz1"
    ws.cell(1, 1, "Hala B0 — elewacje regałów (dane fikcyjne)")
    for ai, aisle in enumerate(aisles):
        ws.cell(3 + ai * band, 1, f"PRZEJŚCIE {aisle}")
    for (r, c), group in cells.items():
        ws.cell(r, c, " | ".join(sorted(group, key=sort_key, reverse=True)))
    wb.save(path)


# --------------------------------------------------------------- próbka EWM
# Szablony gniazd: (pozycji na belce, [(litera, dzielona, typ EWM)]).
def _t(k, spec):
    return k, [(s.rstrip("½").split(":")[0], s.endswith("½"), s.rstrip("½").split(":")[1]) for s in spec.split()]


T = {
    "pal": _t(3, "A:0052 X:0010 Y:0010 Z:0010"),
    "pal_xy": _t(3, "A:0052 X:0010 Y:0010"),
    "pal2": _t(2, "A:0052 X:0010 Y:0010 Z:0010"),
    "pal_b": _t(3, "A:0052 B:0052 X:0010 Y:0010 Z:0010"),
    "dziel": _t(3, "B:0052 C:0010½ D:0010½ X:0010 Y:0010 Z:0010"),
    "dziel2": _t(2, "B:0052 C:0010½ D:0010½ X:0010 Y:0010 Z:0010"),
    "dziel_a": _t(3, "A:0052 C:0010½ D:0010½ X:0010 Y:0010 Z:0010"),
    "pal4": _t(4, "A:0052 X:0010 Y:0010 Z:0010"),
    "gora": _t(3, "A:0010 X:0010 Y:0010 Z:0010"),
    "pal_11": _t(3, "A:0052 X:0011 Y:0011 Z:0011"),
    "pal_50": _t(3, "A:0050 X:0010 Y:0010 Z:0010"),
    "polki": _t(3, "B:0010 C:0010 D:0010 X:0011 Y:0011 Z:0011"),
    "polki_xy": _t(3, "B:0010 C:0010 D:0010 X:0011 Y:0011"),
    "polki_12": _t(3, "B:0012 C:0012 D:0012 X:0012 Y:0012 Z:0012"),
    "xyz": _t(3, "X:0010 Y:0010 Z:0010"),
    "przejazd": _t(4, "Y:0010 Z:0010"),
    "yz2": _t(2, "Y:0010 Z:0010"),
    "jedna": _t(1, "A:0052 X:0010 Y:0010 Z:0010"),
}

# przejście → (numery gniazd, [(szablon, liczba gniazd)]) — pierwszy szablon = domyślny rzędu
# (musi mieć najwięcej gniazd); kolejność gniazd przeplatana deterministycznie.
EWM_AISLES = {
    "07": (list(range(10, 49)), [("pal", 17), ("pal_b", 16), ("pal_xy", 5), ("pal2", 1)]),
    "08": (list(range(10, 48)) + [50], [("pal", 39)]),
    "34": (list(range(10, 52)), [("dziel", 39), ("dziel2", 1), ("dziel_a", 1), ("pal4", 1)]),
    "38": (list(range(10, 29)) + list(range(30, 52)), [("gora", 20), ("pal_11", 11), ("pal_50", 10)]),
    "48": (list(range(10, 46)), [("polki", 23), ("polki_xy", 3), ("polki_12", 10)]),
    "54": ([29, 30] + list(range(33, 65)), [("xyz", 15), ("przejazd", 7), ("yz2", 6), ("jedna", 6)]),
}
# Pojedynczy wyjątek typu EWM (przejście, gniazdo, pozycja, litera) → typ.
EWM_TYPE_EXCEPTIONS = {("08", 22, 1, "A"): "0050"}


def _interleave(assign):
    """[(szablon, n)] → lista szablonów przeplatana (domyślny na pierwszym gnieździe)."""
    pools = [[name] * n for name, n in assign]
    out = []
    while any(pools):
        for pool in pools:
            if pool:
                out.append(pool.pop())
    return out


def ewm_pairs():
    pairs = []
    for aisle, (bays, assign) in EWM_AISLES.items():
        names = _interleave(assign)
        assert len(names) == len(bays), aisle
        for bay, name in zip(bays, names):
            k, levels = T[name]
            for p in range(k):
                for letter, split, typ in levels:
                    typ = EWM_TYPE_EXCEPTIONS.get((aisle, bay, p, letter), typ)
                    for half in ((1, 2) if split else (0,)):
                        pairs.append((f"B0-{aisle}-{bay:02d}{p}{letter}" + (f"-{half}" if half else ""), typ))
    return pairs


def write_ewm_sample(path):
    pairs = ewm_pairs()
    lines = ["# Próbka eksportu EWM — dane syntetyczne (generate_synthetic_b0.py): przejścia B0 "
             + ", ".join(EWM_AISLES), "code;typ"] + [f"{c};{t}" for c, t in pairs]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return len(pairs)


def main():
    codes = hall_codes()
    write_locations_csv(codes, UI_DATA / "warehouse_locations_b0.csv")
    write_map_xlsx(codes, UI_DATA / "warehouse_map_b0.xlsx")
    write_layout_xlsx(codes, UI_DATA / "warehouse_layout_b0.xlsx")
    n_ewm = write_ewm_sample(EWM_SAMPLE)
    aisles = {parts(c)[0] for c in codes}
    print(f"Hala B0 (syntetyczna): {len(codes)} lokalizacji, {len(aisles)} przejść; próbka EWM: {n_ewm} kodów.")


if __name__ == "__main__":
    main()
