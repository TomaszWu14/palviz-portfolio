"""Import eksportu SAP BW „SAP_Dane_materialowe” (.xlsm/.xlsx) do `MaterialMaster`.

Czyta trzy arkusze po nazwie (Nazwy, Hierarchia produktów, Przeliczniki), łączy je po MATNR
i robi upsert (nowe materiały dopisuje, istniejące aktualizuje, nic nie kasuje). Pozostałe
arkusze eksportu (statusy sprzedaży, BLOZ, historia zmian) nie są potrzebne i są pomijane.
Przeliczniki OPZ/KAR/PAZ = liczba jednostek podstawowych (JP) w opakowaniu; OBJ_* w dm³.
"""
SHEETS = ("Nazwy", "Hierarchia produktów", "Przeliczniki")
BATCH = 2000


def norm_matnr(value):
    """„000000000001005796” / 1005796 / „1005796” → „1005796” (klucz jak w zadaniach EWM)."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value or "").strip().lstrip("0")[:18]


def _num(value):
    """Liczba > 0 albo None (SAP wpisuje 0 / puste, gdy opakowania nie ma)."""
    try:
        v = float(str(value).replace(",", ".")) if value not in (None, "") else None
    except ValueError:
        return None
    return v if v and v > 0 else None


def _text(value, n):
    return str(value).strip()[:n] if value is not None else ""


def _sheet_rows(ws):
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
    for r in rows:
        if r and any(v not in (None, "") for v in r):
            yield dict(zip(header, r, strict=False))


def parse_workbook(file):
    """Plik → {matnr: pola MaterialMaster}; ValueError po polsku, gdy brak arkuszy."""
    import openpyxl

    wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    try:
        missing = [s for s in SHEETS if s not in wb.sheetnames]
        if missing:
            raise ValueError("To nie jest eksport „SAP_Dane_materialowe” — brak arkuszy: " + ", ".join(missing))
        out = {}
        for r in _sheet_rows(wb["Nazwy"]):
            m = norm_matnr(r.get("MATNR"))
            if m:
                out[m] = {"ref": _text(r.get("REF"), 60), "name": _text(r.get("TXT_SHORT_PL"), 250),
                          "kind": _text(r.get("RODZAJ"), 4), "hierarchy": _text(r.get("HIERARCHIA"), 18)}
        for r in _sheet_rows(wb["Hierarchia produktów"]):
            m = norm_matnr(r.get("MATNR"))
            if m:
                out.setdefault(m, {}).update({
                    "h1": _text(r.get("H1"), 120), "h2": _text(r.get("H2"), 120),
                    "h3": _text(r.get("H3"), 120), "h4": _text(r.get("H4"), 120),
                    "producer": _text(r.get("PRODUCENT_NAZWA"), 120), "cn_code": _text(r.get("KOD_CN"), 12),
                    "purchase_group": _text(r.get("GRUPA_ZAOPATRZENIOWA"), 10)})
                if not out[m].get("ref"):
                    out[m]["ref"] = _text(r.get("REF"), 60)
        for r in _sheet_rows(wb["Przeliczniki"]):
            m = norm_matnr(r.get("MATNR"))
            if m:
                out.setdefault(m, {}).update({
                    "base_unit": _text(r.get("JP"), 6),
                    "pcs_per_opz": _num(r.get("OPZ")), "pcs_per_carton": _num(r.get("KAR")),
                    "pcs_per_pallet": _num(r.get("PAZ")),
                    "vol_unit_dm3": _num(r.get("OBJ_SZT_OP")), "vol_carton_dm3": _num(r.get("OBJ_KAR")),
                    "vol_pallet_dm3": _num(r.get("OBJ_PAZ")),
                    "weight_unit_kg": _num(r.get("WAGA_SZT_OP")), "weight_carton_kg": _num(r.get("WAGA_KAR")),
                    "weight_pallet_kg": _num(r.get("WAGA_PAZ"))})
                if not out[m].get("ref"):
                    out[m]["ref"] = _text(r.get("REF"), 60)
        return out
    finally:
        wb.close()


def upsert(materials):
    """Zapis partiami (INSERT … ON CONFLICT (matnr) DO UPDATE) → (nowe, zaktualizowane)."""
    from .models import MaterialMaster

    fields = [f.name for f in MaterialMaster._meta.concrete_fields if f.name not in ("id", "matnr")]
    existing = set(MaterialMaster.objects.filter(matnr__in=list(materials)).values_list("matnr", flat=True))
    objs = [MaterialMaster(matnr=m, **vals) for m, vals in materials.items()]
    MaterialMaster.objects.bulk_create(objs, batch_size=BATCH, update_conflicts=True,
                                       unique_fields=["matnr"], update_fields=fields)
    return len(materials) - len(existing), len(existing)


def stats(materials):
    return {"total": len(materials),
            "with_hierarchy": sum(1 for v in materials.values() if v.get("h1")),
            "with_pallet": sum(1 for v in materials.values() if v.get("pcs_per_pallet")),
            "with_carton": sum(1 for v in materials.values() if v.get("pcs_per_carton"))}
