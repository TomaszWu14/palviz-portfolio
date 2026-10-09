"""Klasy lokalizacji (symulacja) — CRUD, symulacja, fit-all, szablon/import xlsx (ui/views/locations.py, TEST-004)."""

import io

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from ui.models import (
    Carton,
    WarehouseLocationMaster,
    WarehouseLocationMasterBatch,
    WarehouseLocationType,
)
from ui.roles import GROUP_MASTER_DATA, GROUP_VIEWER


def _user(username, group_name=None):
    u = get_user_model().objects.create_user(username=username, password="x")
    if group_name:
        u.groups.add(Group.objects.get_or_create(name=group_name)[0])
    return u


def _loc(name="Regał A", **kw):
    data = {"name": name, "width_cm": 90, "depth_cm": 120, "total_height_cm": 200}
    data.update(kw)
    return WarehouseLocationType.objects.create(**data)


LOC_POST = {
    "name": "Nowy regał",
    "location_class": "pallet_full",
    "is_pallet_location": "on",
    "width_cm": 90,
    "depth_cm": 120,
    "total_height_cm": 200,
    "pallet_height_cm": 15,
    "manipulation_margin_cm": 20,
    "is_active": "on",
}
SIM_GET = {"carton_l": 40, "carton_w": 30, "carton_h": 25, "unit_weight": "0.45", "pcs_per_carton": 24}


class LocationCrudTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.md = _user("md", GROUP_MASTER_DATA)
        cls.viewer = _user("viewer", GROUP_VIEWER)

    def setUp(self):
        self.client.force_login(self.md)

    def test_list(self):
        _loc()
        r = self.client.get(reverse("ui:planner_locations"))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(len(r.context["locations"]), 1)

    def test_list_forbidden_without_role(self):
        self.client.force_login(_user("norole"))
        self.assertEqual(self.client.get(reverse("ui:planner_locations")).status_code, 403)

    def test_viewer_cannot_write(self):
        loc = _loc()
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get(reverse("ui:planner_location_new")).status_code, 403)
        self.assertEqual(self.client.post(reverse("ui:planner_location_new"), LOC_POST).status_code, 403)
        self.assertEqual(self.client.post(reverse("ui:planner_location_delete", args=[loc.pk])).status_code, 403)
        self.assertEqual(WarehouseLocationType.objects.count(), 1)

    def test_get_new_and_edit_forms(self):
        loc = _loc()
        self.assertEqual(self.client.get(reverse("ui:planner_location_new")).status_code, 200)
        r = self.client.get(reverse("ui:planner_location_edit", args=[loc.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context["instance"], loc)

    def test_create(self):
        r = self.client.post(reverse("ui:planner_location_new"), LOC_POST)
        self.assertRedirects(r, reverse("ui:planner_locations"), fetch_redirect_response=False)
        self.assertTrue(WarehouseLocationType.objects.filter(name="Nowy regał", width_cm=90).exists())

    def test_edit(self):
        loc = _loc()
        r = self.client.post(reverse("ui:planner_location_edit", args=[loc.pk]), {**LOC_POST, "width_cm": 100})
        self.assertEqual(r.status_code, 302)
        loc.refresh_from_db()
        self.assertEqual((loc.name, loc.width_cm), ("Nowy regał", 100))

    def test_validation_error(self):
        r = self.client.post(reverse("ui:planner_location_new"), {**LOC_POST, "width_cm": ""})
        self.assertEqual(r.status_code, 200)
        self.assertIn("width_cm", r.context["form"].errors)
        self.assertFalse(WarehouseLocationType.objects.exists())

    def test_delete_confirm_then_post(self):
        loc = _loc()
        r = self.client.get(reverse("ui:planner_location_delete", args=[loc.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context["loc"], loc)
        r = self.client.post(reverse("ui:planner_location_delete", args=[loc.pk]))
        self.assertRedirects(r, reverse("ui:planner_locations"), fetch_redirect_response=False)
        self.assertFalse(WarehouseLocationType.objects.exists())

    def test_missing_pk_404(self):
        self.assertEqual(self.client.get(reverse("ui:planner_location_edit", args=[999])).status_code, 404)


class LocationSimulateTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.viewer = _user("viewer", GROUP_VIEWER)
        cls.loc = _loc(is_pallet_location=True)
        cls.shelf = _loc("Półka", is_pallet_location=False, location_class="shelf")
        cls.carton = Carton.objects.create(
            name="K", length_cm=40, width_cm=30, height_cm=25, unit_weight_kg=0.5, pieces_per_carton=10
        )

    def setUp(self):
        self.client.force_login(self.viewer)

    def test_simulate_empty_form(self):
        r = self.client.get(reverse("ui:planner_location_simulate", args=[self.loc.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertIsNone(r.context["fig_json"])

    def test_simulate_manual_dimensions_uses_location_pallet_flag(self):
        r = self.client.get(reverse("ui:planner_location_simulate", args=[self.shelf.pk]), SIM_GET)
        self.assertEqual(r.status_code, 200)
        self.assertIsNotNone(r.context["fig_json"])
        self.assertIsNotNone(r.context["fig_2d_json"])
        self.assertFalse(r.context["with_pallet"])

    def test_simulate_carton_from_db_explicit_no_pallet(self):
        url = reverse("ui:planner_location_simulate", args=[self.loc.pk])
        r = self.client.get(url, {"carton_from_db": self.carton.pk, "with_pallet": "0"})
        self.assertEqual(r.status_code, 200)
        self.assertIsNotNone(r.context["fit"])
        self.assertFalse(r.context["with_pallet"])

    def test_simulate_invalid_form_shows_errors(self):
        r = self.client.get(reverse("ui:planner_location_simulate", args=[self.loc.pk]), {"carton_l": 40})
        self.assertEqual(r.status_code, 200)
        self.assertIn("carton_w", r.context["form"].errors)
        self.assertIsNone(r.context["fig_json"])

    def test_fit_all_empty(self):
        r = self.client.get(reverse("ui:planner_location_fit_all"))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context["results"], [])

    def test_fit_all_manual_and_db_carton(self):
        r = self.client.get(reverse("ui:planner_location_fit_all"), SIM_GET)
        self.assertEqual(r.status_code, 200)
        results = r.context["results"]
        self.assertEqual(len(results), 2)
        by_name = {row["loc"].name: row for row in results}
        # tryb główny wynika z typu lokalizacji, drugi to widok alternatywny
        self.assertEqual([m["with_pallet"] for m in by_name["Regał A"]["modes"]], [True, False])
        self.assertEqual([m["with_pallet"] for m in by_name["Półka"]["modes"]], [False, True])
        r = self.client.get(reverse("ui:planner_location_fit_all"), {"carton_from_db": self.carton.pk})
        self.assertEqual(r.context["cl"], 40)

    def test_fit_all_forbidden_without_role(self):
        self.client.force_login(_user("norole"))
        self.assertEqual(self.client.get(reverse("ui:planner_location_fit_all")).status_code, 403)


def _xlsx(rows, description_row=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["location_code", "warehouse_type", "capacity_mm"])
    if description_row:
        ws.append(["Kod", "Typ", "Pojemność"])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(
        "lok.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


class LocationExcelTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.md = _user("md", GROUP_MASTER_DATA)
        cls.viewer = _user("viewer", GROUP_VIEWER)

    def setUp(self):
        self.client.force_login(self.md)

    def _msgs(self, r):
        return [str(m) for m in r.context["messages"]]

    def test_template_download(self):
        r = self.client.get(reverse("ui:excel_template_locations"))
        self.assertEqual(r.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Lokalizacje", "Format kodu"])
        self.assertEqual(wb["Lokalizacje"]["A1"].value, "location_code")

    def test_import_get_redirects(self):
        r = self.client.get(reverse("ui:excel_import_locations"))
        self.assertRedirects(r, reverse("ui:planner_excel_templates"), fetch_redirect_response=False)

    def test_import_no_file(self):
        r = self.client.post(reverse("ui:excel_import_locations"))
        self.assertRedirects(r, reverse("ui:planner_excel_templates"), fetch_redirect_response=False)

    def test_import_viewer_forbidden(self):
        self.client.force_login(self.viewer)
        r = self.client.post(reverse("ui:excel_import_locations"), {"file": _xlsx([["B0-01-300A", "B0", 2100]])})
        self.assertEqual(r.status_code, 403)
        self.assertFalse(WarehouseLocationMaster.objects.exists())

    def test_import_creates_new_active_batch(self):
        old = WarehouseLocationMasterBatch.objects.create(name="stary", is_active=True)
        rows = [
            ["B0-01-300A", "B0", 2100],
            ["", "B0", 100],  # brak kodu → błąd
            ["B0-01-300B", "B0", ""],  # pusta pojemność → 0
        ]
        r = self.client.post(reverse("ui:excel_import_locations"), {"file": _xlsx(rows)}, follow=True)
        old.refresh_from_db()
        self.assertFalse(old.is_active)
        batch = WarehouseLocationMasterBatch.objects.get(is_active=True)
        self.assertEqual(batch.location_count, 2)
        heights = dict(batch.locations.values_list("location_code", "height_mm"))
        self.assertEqual(heights, {"B0-01-300A": 2100, "B0-01-300B": 0})
        self.assertTrue(any("2 wpisów, 1 błędów" in m for m in self._msgs(r)))

    def test_import_without_errors_success_message(self):
        r = self.client.post(
            reverse("ui:excel_import_locations"),
            {"file": _xlsx([["C1-01-200A", "C1", 2200]], description_row=False)},
            follow=True,
        )
        self.assertTrue(any("1 wpisów, 0 błędów" in m for m in self._msgs(r)))

    def test_import_broken_file(self):
        bad = SimpleUploadedFile("x.xlsx", b"not a zip", content_type="application/octet-stream")
        r = self.client.post(reverse("ui:excel_import_locations"), {"file": bad}, follow=True)
        self.assertTrue(any("Błąd wczytywania pliku" in m for m in self._msgs(r)))
        self.assertFalse(WarehouseLocationMasterBatch.objects.exists())

    def test_import_out_of_range_capacity_rejects_row_not_whole_import(self):
        """Regresja: pojemność poza int4 / inf / NaN / tekst → wcześniej None w NOT NULL
        height_mm (IntegrityError) albo OverflowError → przerwany CAŁY import. Teraz taki
        wiersz jest odrzucany z komunikatem, a poprawne wiersze trafiają do batcha."""
        rows = [
            ["B0-01-300A", "B0", 2100],
            ["B0-01-300B", "B0", 1e30],        # ogromna pojemność (poza int4)
            ["B0-01-300C", "B0", "inf"],
            ["B0-01-300D", "B0", "nan"],
            ["B0-01-300E", "B0", "dużo"],       # tekst nieparsowalny
            ["B0-01-300F", "B0", 9999999999],   # > 2 147 483 647
            ["B0-01-300G", "B0", 1800],
        ]
        r = self.client.post(reverse("ui:excel_import_locations"), {"file": _xlsx(rows)}, follow=True)
        self.assertEqual(r.status_code, 200)
        batch = WarehouseLocationMasterBatch.objects.get(is_active=True)
        heights = dict(batch.locations.values_list("location_code", "height_mm"))
        self.assertEqual(heights, {"B0-01-300A": 2100, "B0-01-300G": 1800})
        self.assertEqual(batch.location_count, 2)
        msgs = self._msgs(r)
        self.assertFalse(any("Błąd wczytywania pliku" in m for m in msgs), msgs)
        self.assertTrue(any("2 wpisów, 5 błędów" in m for m in msgs), msgs)
        details = " ".join(msgs)
        self.assertIn("capacity_mm", details)
        self.assertIn("B0-01-300E", details)
        self.assertIn("dużo", details)
        self.assertIn("wiersz 5", details)   # nagłówek + opis + 3 wiersze → B0-01-300C to wiersz 5
