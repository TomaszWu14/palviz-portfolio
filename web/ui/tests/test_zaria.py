"""ZARIA (internal AI chat) module: module gating, sending a message (mocked LLM
call), role×model access grid, rate limiting, and the budget-alert dedup."""
import importlib.util
import json
import unittest
from types import SimpleNamespace
from unittest.mock import patch

_HAS_OPENAI = importlib.util.find_spec("openai") is not None

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.cache import cache
from django.test import SimpleTestCase, TestCase
from django.urls import reverse

from ui.models import (ZariaModel, ZariaModelRoleAccess, ZariaUserModelAccess, ZariaConfig,
                       ZariaRoleTokenBudget, ZariaConversation, ZariaMessage, Task)
from ui.roles import GROUP_ADMIN, GROUP_MASTER_DATA, GROUP_CLIENT, GROUP_TRANSPORT
from ui.views.zaria import _check_budget_threshold, can_use_zaria_model
from ui.zaria_llm import _complete_openai


@unittest.skipUnless(_HAS_OPENAI, "openai SDK opcjonalne — pomiń gdy niezainstalowane")
class CompleteOpenAIUsageTests(SimpleTestCase):
    """Ollama bywa bez `usage` — poprawna odpowiedź nie może przepaść (jak w streamie)."""
    @patch("openai.OpenAI")
    def test_missing_usage_returns_zero_tokens_not_crash(self, mock_openai):
        resp = SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content="Odpowiedź Ollamy"))],
            usage=None)
        mock_openai.return_value.chat.completions.create.return_value = resp
        model = SimpleNamespace(provider="ollama", key="llama3")
        text, pt, ct = _complete_openai(model, [{"role": "user", "content": "hej"}], "", 100)
        self.assertEqual(text, "Odpowiedź Ollamy")
        self.assertEqual((pt, ct), (0, 0))

    @patch("openai.OpenAI")
    def test_empty_choices_no_index_error(self, mock_openai):
        mock_openai.return_value.chat.completions.create.return_value = SimpleNamespace(
            choices=[], usage=None)
        model = SimpleNamespace(provider="ollama", key="llama3")
        self.assertEqual(_complete_openai(model, [{"role": "user", "content": "x"}], "", 100),
                         ("", 0, 0))


def _user(name, *groups):
    u = get_user_model().objects.create_user(username=name, password="x")
    for g in groups:
        u.groups.add(Group.objects.get_or_create(name=g)[0])
    return u


class ZariaModuleAccessTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = _user("zadm", GROUP_ADMIN)
        cls.md = _user("zmd", GROUP_MASTER_DATA)
        cls.klient = _user("zcli", GROUP_CLIENT)

    def setUp(self):
        cache.clear()

    def test_admin_sees_zaria_home(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("ui:zaria_home")).status_code, 200)

    def test_internal_role_sees_zaria_home(self):
        # Rollout (P7): role wewnętrzne (tu Master Data) widzą moduł — czat działa po
        # nadaniu modelu, ale sam hub/home jest dostępny.
        self.client.force_login(self.md)
        self.assertEqual(self.client.get(reverse("ui:zaria_home")).status_code, 200)

    def test_client_role_gets_403(self):
        # Obsługa klienta (zewnętrzna) jako jedyna wewnętrzna-wykluczona nie ma dostępu.
        self.client.force_login(self.klient)
        self.assertEqual(self.client.get(reverse("ui:zaria_home")).status_code, 403)


class ZariaSendMessageTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = _user("zadm2", GROUP_ADMIN)
        cls.model = ZariaModel.objects.create(
            key="claude-sonnet-4-5", display_name="Claude Sonnet", provider="anthropic",
            price_input_per_1k="0.01", price_output_per_1k="0.03")
        ZariaModelRoleAccess.objects.create(model=cls.model, group_name=GROUP_ADMIN)

    def setUp(self):
        cache.clear()
        self.client.force_login(self.admin)
        self.conv = ZariaConversation.objects.create(user=self.admin, model=self.model)

    @patch("ui.zaria_llm.complete")
    def test_send_message_logs_cost_and_tokens(self, mock_complete):
        mock_complete.return_value = ("Cześć!", 100, 50, 250)
        r = self.client.post(reverse("ui:zaria_send_message", args=[self.conv.pk]),
                             {"content": "Witaj"})
        self.assertRedirects(r, reverse("ui:zaria_conversation", args=[self.conv.pk]))

        msgs = list(self.conv.messages.all())
        self.assertEqual(len(msgs), 2)
        user_msg, assistant_msg = msgs
        self.assertEqual(user_msg.role, "user")
        self.assertEqual(assistant_msg.role, "assistant")
        self.assertEqual(assistant_msg.prompt_tokens, 100)
        self.assertEqual(assistant_msg.completion_tokens, 50)
        # 100/1000*0.01 + 50/1000*0.03 = 0.001 + 0.0015 = 0.0025
        self.assertAlmostEqual(float(assistant_msg.cost_pln), 0.0025)

    @patch("ui.zaria_llm.complete")
    def test_vendor_error_is_recorded_not_500(self, mock_complete):
        from ui.zaria_llm import ZariaLLMError
        mock_complete.side_effect = ZariaLLMError("boom")
        r = self.client.post(reverse("ui:zaria_send_message", args=[self.conv.pk]),
                             {"content": "Witaj"})
        self.assertEqual(r.status_code, 302)
        assistant_msg = self.conv.messages.filter(role="assistant").first()
        self.assertTrue(assistant_msg.error)
        self.assertEqual(assistant_msg.content, "")

    def test_no_access_to_model_blocks_send(self):
        # A model nobody has been granted access to (not even GROUP_ADMIN) — the
        # per-model grid is a finer-grained check than the module-level role gate.
        ungranted_model = ZariaModel.objects.create(
            key="claude-opus", display_name="Claude Opus", provider="anthropic")
        conv = ZariaConversation.objects.create(user=self.admin, model=ungranted_model)
        r = self.client.post(reverse("ui:zaria_send_message", args=[conv.pk]), {"content": "hej"})
        self.assertRedirects(r, reverse("ui:zaria_conversation", args=[conv.pk]))
        self.assertEqual(conv.messages.count(), 0)


class ZariaEffortTests(TestCase):
    """Konfigurowalny wysiłek (effort) rozmowy: zapis przy starcie + przekazanie do LLM."""
    @classmethod
    def setUpTestData(cls):
        cls.admin = _user("zeff", GROUP_ADMIN)
        cls.model = ZariaModel.objects.create(
            key="claude-opus-4-8", display_name="Opus", provider="anthropic",
            price_input_per_1k="0.02", price_output_per_1k="0.10")
        ZariaModelRoleAccess.objects.create(model=cls.model, group_name=GROUP_ADMIN)

    def setUp(self):
        cache.clear()
        self.client.force_login(self.admin)

    def test_new_conversation_saves_effort(self):
        r = self.client.post(reverse("ui:zaria_new_conversation"),
                             {"model": str(self.model.id), "effort": "high"})
        conv = ZariaConversation.objects.get(user=self.admin)
        self.assertEqual(r.status_code, 302)
        self.assertEqual(conv.effort, "high")

    def test_invalid_effort_falls_back_to_auto(self):
        self.client.post(reverse("ui:zaria_new_conversation"),
                         {"model": str(self.model.id), "effort": "xxl"})
        self.assertEqual(ZariaConversation.objects.get(user=self.admin).effort, "auto")

    @patch("ui.zaria_llm.complete")
    def test_effort_passed_to_llm_call(self, mock_complete):
        mock_complete.return_value = ("ok", 1, 1, 10)
        conv = ZariaConversation.objects.create(user=self.admin, model=self.model, effort="medium")
        self.client.post(reverse("ui:zaria_send_message", args=[conv.pk]), {"content": "hej"})
        self.assertEqual(mock_complete.call_args.kwargs.get("effort"), "medium")

    @patch("ui.zaria_llm.complete")
    def test_api_chat_passes_effort(self, mock_complete):
        mock_complete.return_value = ("ok", 1, 1, 10)
        r = self.client.post(reverse("ui:zaria_api_chat"),
                             json.dumps({"message": "hej", "effort": "high"}),
                             content_type="application/json")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(mock_complete.call_args.kwargs.get("effort"), "high")

    @patch("ui.zaria_llm.complete")
    def test_high_effort_denied_for_non_privileged_role(self, mock_complete):
        # INT-007: 'high' (sufit 16000 tokenów) tylko Administratorzy/Master Data.
        mock_complete.return_value = ("ok", 1, 1, 10)
        ZariaModelRoleAccess.objects.create(model=self.model, group_name=GROUP_TRANSPORT)
        user = _user("zeff_tr", GROUP_TRANSPORT)
        self.client.force_login(user)
        self.client.post(reverse("ui:zaria_new_conversation"), {"model": str(self.model.id), "effort": "high"})
        self.assertEqual(ZariaConversation.objects.get(user=user).effort, "auto")
        self.client.post(reverse("ui:zaria_api_chat"), json.dumps({"message": "hej", "effort": "high"}),
                         content_type="application/json")
        self.assertEqual(mock_complete.call_args.kwargs.get("effort"), "auto")

    def test_thinking_params_mapping(self):
        from ui.zaria_llm import _thinking_params
        extra, mt = _thinking_params(self.model, "high", 1000)
        self.assertEqual(extra["thinking"], {"type": "adaptive"})
        self.assertEqual(extra["output_config"], {"effort": "high"})
        self.assertEqual(mt, 16000)   # sufit podniesiony — myślenie liczy się do max_tokens
        # auto → bez zmian
        self.assertEqual(_thinking_params(self.model, "auto", 1000), ({}, 1000))
        # Haiku nie wspiera effort — pomijamy parametry
        haiku = ZariaModel(key="claude-haiku-4-5", provider="anthropic")
        self.assertEqual(_thinking_params(haiku, "high", 1000), ({}, 1000))


class ZariaMsgXlsxTests(TestCase):
    """Eksport tabeli markdown z wiadomości asystenta do .xlsx."""
    TABLE = ("| Lp | Produkt | Ilość |\n"
             "|---|---|---|\n"
             "| 1 | Strzykawka 2ml | 5 |\n"
             "| 2 | Cewnik Foley | 20 |\n")

    @classmethod
    def setUpTestData(cls):
        cls.user = _user("zxlsx", GROUP_ADMIN)
        cls.other = _user("zxlsx2", GROUP_ADMIN)
        cls.model = ZariaModel.objects.create(key="claude-opus-4-8", display_name="Opus",
                                              provider="anthropic")
        cls.conv = ZariaConversation.objects.create(user=cls.user, model=cls.model)
        cls.msg = ZariaMessage.objects.create(conversation=cls.conv, role="assistant",
                                              content="Oto dane:\n\n" + cls.TABLE)

    def setUp(self):
        cache.clear()

    def test_has_table_and_button_rendered(self):
        self.assertTrue(self.msg.has_table)
        self.client.force_login(self.user)
        r = self.client.get(reverse("ui:zaria_conversation", args=[self.conv.pk]))
        self.assertContains(r, reverse("ui:zaria_msg_xlsx", args=[self.msg.pk]))

    def test_xlsx_download_contains_rows(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.client.force_login(self.user)
        r = self.client.get(reverse("ui:zaria_msg_xlsx", args=[self.msg.pk]))
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(BytesIO(r.content)).active
        self.assertEqual(ws["B1"].value, "Produkt")
        self.assertEqual(ws["C3"].value, 20)          # liczby jako liczby
        self.assertEqual(ws.max_row, 3)               # separator pominięty

    def test_other_user_gets_404(self):
        self.client.force_login(self.other)
        r = self.client.get(reverse("ui:zaria_msg_xlsx", args=[self.msg.pk]))
        self.assertEqual(r.status_code, 404)

    def test_message_without_table_redirects(self):
        msg = ZariaMessage.objects.create(conversation=self.conv, role="assistant", content="Cześć")
        self.assertFalse(msg.has_table)
        self.client.force_login(self.user)
        r = self.client.get(reverse("ui:zaria_msg_xlsx", args=[msg.pk]))
        self.assertRedirects(r, reverse("ui:zaria_conversation", args=[self.conv.pk]))


class ZariaRoleAccessGridTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = _user("zadm3", GROUP_ADMIN)
        cls.md_user = _user("zmd3", GROUP_MASTER_DATA)
        cls.model = ZariaModel.objects.create(
            key="claude-haiku", display_name="Claude Haiku", provider="anthropic")

    def test_grid_post_creates_role_access_rows(self):
        self.client.force_login(self.admin)
        r = self.client.post(reverse("ui:admin_zaria_role_access"), {
            f"a_{self.model.pk}_{GROUP_MASTER_DATA}": "1",
        })
        self.assertRedirects(r, reverse("ui:admin_zaria_role_access"))
        self.assertTrue(ZariaModelRoleAccess.objects.filter(
            model=self.model, group_name=GROUP_MASTER_DATA).exists())
        self.assertTrue(can_use_zaria_model(self.md_user, self.model))

    def test_grid_rebuild_removes_unticked_rows(self):
        ZariaModelRoleAccess.objects.create(model=self.model, group_name=GROUP_MASTER_DATA)
        self.client.force_login(self.admin)
        self.client.post(reverse("ui:admin_zaria_role_access"), {})   # nothing ticked
        self.assertFalse(ZariaModelRoleAccess.objects.filter(model=self.model).exists())

    def test_user_override_wins_over_role_grid(self):
        ZariaModelRoleAccess.objects.create(model=self.model, group_name=GROUP_MASTER_DATA)
        ZariaUserModelAccess.objects.create(user=self.md_user, model=self.model, allowed=False)
        self.assertFalse(can_use_zaria_model(self.md_user, self.model))


class ZariaRateLimitTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = _user("zadm4", GROUP_ADMIN)
        cls.model = ZariaModel.objects.create(
            key="claude-sonnet-4-5", display_name="Claude Sonnet", provider="anthropic")
        ZariaModelRoleAccess.objects.create(model=cls.model, group_name=GROUP_ADMIN)

    def setUp(self):
        cache.clear()
        config = ZariaConfig.load()
        config.rate_limit_per_minute = 1
        config.rate_limit_per_day = 200
        config.save()
        self.client.force_login(self.admin)
        self.conv = ZariaConversation.objects.create(user=self.admin, model=self.model)

    @patch("ui.zaria_llm.complete")
    def test_second_message_within_a_minute_is_blocked(self, mock_complete):
        mock_complete.return_value = ("ok", 1, 1, 10)
        self.client.post(reverse("ui:zaria_send_message", args=[self.conv.pk]), {"content": "1"})
        self.assertEqual(mock_complete.call_count, 1)
        self.client.post(reverse("ui:zaria_send_message", args=[self.conv.pk]), {"content": "2"})
        self.assertEqual(mock_complete.call_count, 1)   # second call rate-limited, not forwarded


class ZariaBudgetAlertTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = _user("zbudget", GROUP_ADMIN)

    def test_threshold_alert_is_deduped_per_month(self):
        _check_budget_threshold(self.user, 150, 100)
        _check_budget_threshold(self.user, 160, 100)
        self.assertEqual(Task.objects.filter(category="zaria_budget").count(), 1)

    def test_no_alert_under_budget(self):
        _check_budget_threshold(self.user, 50, 100)
        self.assertEqual(Task.objects.filter(category="zaria_budget").count(), 0)


class ZariaTokenBudgetTests(TestCase):
    """Egzekwowanie miesięcznego limitu TOKENÓW (obok budżetu PLN) + per-rola."""
    @classmethod
    def setUpTestData(cls):
        cls.user = _user("ztok", GROUP_ADMIN)
        cls.model = ZariaModel.objects.create(
            key="claude-haiku", display_name="Haiku", provider="anthropic",
            price_input_per_1k="0.01", price_output_per_1k="0.03")
        ZariaModelRoleAccess.objects.create(model=cls.model, group_name=GROUP_ADMIN)

    def setUp(self):
        cache.clear()
        self.client.force_login(self.user)
        self.conv = ZariaConversation.objects.create(user=self.user, model=self.model)

    def _seed_tokens(self, tokens):
        ZariaMessage.objects.create(conversation=self.conv, role="assistant", model=self.model,
                                    prompt_tokens=tokens, completion_tokens=0, cost_pln=0)

    def test_default_token_budget_blocks_html_send(self):
        cfg = ZariaConfig.load(); cfg.default_monthly_token_budget = 1000; cfg.save()
        self._seed_tokens(1000)   # zużyto = limit → blokada
        r = self.client.post(reverse("ui:zaria_send_message", args=[self.conv.pk]), {"content": "hej"})
        self.assertRedirects(r, reverse("ui:zaria_conversation", args=[self.conv.pk]))
        self.assertEqual(self.conv.messages.filter(role="user").count(), 0)  # nic nie wysłano

    def test_zero_budget_means_unlimited(self):
        from ui.views.zaria import token_budget_for
        cfg = ZariaConfig.load(); cfg.default_monthly_token_budget = 0; cfg.save()
        self._seed_tokens(10_000_000)
        self.assertEqual(token_budget_for(self.user), 0)

    def test_role_budget_overrides_default(self):
        from ui.views.zaria import token_budget_for
        cfg = ZariaConfig.load(); cfg.default_monthly_token_budget = 100; cfg.save()
        ZariaRoleTokenBudget.objects.create(group_name=GROUP_ADMIN, monthly_tokens=9_000_000)
        self.assertEqual(token_budget_for(self.user), 9_000_000)
