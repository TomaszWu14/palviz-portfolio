"""Import użytkowników z XLSX + brakujące wzory plików importu.

Bezpieczeństwo: hasło z pliku jest hashowane, konto dostaje flagę wymuszonej zmiany,
a middleware nie wpuszcza takiego konta dalej niż ekran zmiany hasła."""
import io

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook

from ui.models import UserProfile, ControllerZone, UserModuleAccess
from ui.roles import GROUP_ADMIN, GROUP_CONTROLLER, GROUP_MASTER_DATA, ALL_GROUPS

HEADERS = ["username", "first_name", "last_name", "email", "password",
           "must_change_password", "is_active", "is_staff", "is_superuser",
           "groups", "phone", "department", "section",
           "control_zones", "modules_allow", "modules_deny"]


def _xlsx(rows):
    """Zbuduj plik w formacie wzoru: 1=nagłówki, 2=opisy, 3+=dane."""
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    ws.append(["opis"] * len(HEADERS))
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "users.xlsx"
    return buf


def _admin(name="imp_admin"):
    u = get_user_model().objects.create_user(username=name, password="x", is_superuser=True)
    u.groups.add(Group.objects.get_or_create(name=GROUP_ADMIN)[0])
    return u


class UserImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        for g in ALL_GROUPS:
            Group.objects.get_or_create(name=g)
        cls.admin = _admin()

    def setUp(self):
        self.client.force_login(self.admin)

    def _post(self, rows):
        return self.client.post(reverse("ui:excel_import_users"),
                                {"file": _xlsx(rows)}, follow=True)

    def test_creates_user_with_roles_profile_zones_and_modules(self):
        self._post([["jkowalski", "Jan", "Kowalski", "j@z.pl", "", "", 1, 0, 0,
                     f"{GROUP_CONTROLLER}; {GROUP_MASTER_DATA}", "600100200", "Magazyn", "GLS",
                     "92JU; 92EX", "kontrola_hu", "paletyzacja"]])
        u = get_user_model().objects.get(username="jkowalski")
        self.assertEqual((u.first_name, u.last_name, u.email), ("Jan", "Kowalski", "j@z.pl"))
        self.assertEqual(set(u.groups.values_list("name", flat=True)),
                         {GROUP_CONTROLLER, GROUP_MASTER_DATA})
        prof = UserProfile.objects.get(user=u)
        self.assertEqual((prof.phone, prof.department, prof.section),
                         ("600100200", "Magazyn", "GLS"))
        self.assertEqual(set(ControllerZone.objects.filter(user=u).values_list("code", flat=True)),
                         {"92JU", "92EX"})
        self.assertTrue(UserModuleAccess.objects.get(user=u, module_key="kontrola_hu").allowed)
        self.assertFalse(UserModuleAccess.objects.get(user=u, module_key="paletyzacja").allowed)

    def test_password_is_hashed_and_forces_change(self):
        self._post([["haslowy", "", "", "", "Tajne!2026", "", 1, 0, 0, "", "", "", "", "", "", ""]])
        u = get_user_model().objects.get(username="haslowy")
        self.assertNotEqual(u.password, "Tajne!2026")      # nigdy plaintext w bazie
        self.assertTrue(u.check_password("Tajne!2026"))
        self.assertTrue(u.profile.must_change_password)     # domyślnie wymuszone

    def test_no_password_means_unusable(self):
        self._post([["bezhasla", "", "", "", "", "", 1, 0, 0, "", "", "", "", "", "", ""]])
        u = get_user_model().objects.get(username="bezhasla")
        self.assertFalse(u.has_usable_password())
        self.assertFalse(u.profile.must_change_password)

    def test_upsert_does_not_duplicate(self):
        row = ["dwarazy", "Jan", "", "", "", "", 1, 0, 0, "", "", "", "", "", "", ""]
        self._post([row])
        self._post([["dwarazy", "Janusz", "", "", "", "", 1, 0, 0, "", "", "", "", "", "", ""]])
        users = get_user_model().objects.filter(username="dwarazy")
        self.assertEqual(users.count(), 1)
        self.assertEqual(users.first().first_name, "Janusz")

    def test_row_without_username_is_reported_not_fatal(self):
        r = self._post([["", "Bez", "Loginu", "", "", "", 1, 0, 0, "", "", "", "", "", "", ""],
                        ["zloginem", "", "", "", "", "", 1, 0, 0, "", "", "", "", "", "", ""]])
        self.assertTrue(get_user_model().objects.filter(username="zloginem").exists())
        self.assertContains(r, "brak username")

    def test_unknown_group_and_module_are_ignored(self):
        self._post([["czysty", "", "", "", "", "", 1, 0, 0, "Nieistniejąca Rola", "", "", "",
                     "", "nieistniejacy_modul", ""]])
        u = get_user_model().objects.get(username="czysty")
        self.assertEqual(u.groups.count(), 0)
        self.assertEqual(UserModuleAccess.objects.filter(user=u).count(), 0)

    def test_non_admin_cannot_import(self):
        other = get_user_model().objects.create_user("zwykly", password="x")
        self.client.force_login(other)
        r = self.client.post(reverse("ui:excel_import_users"), {"file": _xlsx([])})
        self.assertEqual(r.status_code, 403)


class PasswordChangeRequiredTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("musizmienic", password="Stare!2026")
        prof, _ = UserProfile.objects.get_or_create(user=self.user)
        prof.must_change_password = True
        prof.save()
        self.client.force_login(self.user)

    def test_redirected_until_password_changed(self):
        r = self.client.get(reverse("ui:home"))
        self.assertRedirects(r, reverse("ui:password_change"))

    def test_change_clears_flag_and_unblocks(self):
        self.client.post(reverse("ui:password_change"),
                         {"old_password": "Stare!2026",
                          "new_password1": "ZupelnieNowe!2026",
                          "new_password2": "ZupelnieNowe!2026"})
        self.user.profile.refresh_from_db()
        self.assertFalse(self.user.profile.must_change_password)
        # Middleware już nie zawraca na zmianę hasła (dokąd trafi hub, zależy od ról).
        r = self.client.get(reverse("ui:home"))
        self.assertNotEqual(r.get("Location"), reverse("ui:password_change"))


class ImportTemplateTests(TestCase):
    """Każdy wiersz centrum importu ma wzór do pobrania (XLSX, poprawny plik)."""
    @classmethod
    def setUpTestData(cls):
        cls.admin = _admin("tpl_admin")

    def setUp(self):
        self.client.force_login(self.admin)

    def test_all_templates_download(self):
        from openpyxl import load_workbook
        for name in ("excel_template_customers", "excel_template_stock_hu",
                     "excel_template_marm", "excel_template_users", "excel_template_fix"):
            with self.subTest(name=name):
                r = self.client.get(reverse(f"ui:{name}"))
                self.assertEqual(r.status_code, 200)
                self.assertIn("spreadsheetml", r["Content-Type"])
                ws = load_workbook(io.BytesIO(r.content)).active
                self.assertGreaterEqual(ws.max_row, 3)      # nagłówki + opisy + przykłady

    def test_data_center_shows_template_for_every_import(self):
        r = self.client.get(reverse("ui:data_center"))
        for row in r.context["imports"]:
            with self.subTest(label=row["label"]):
                self.assertIsNotNone(row["tpl"], f"brak wzoru dla: {row['label']}")
        labels = [row["label"] for row in r.context["imports"]]
        self.assertTrue(any("Użytkownicy" in l for l in labels))
