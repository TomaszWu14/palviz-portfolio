# Auto-split from the former monolithic views.py. Feature views live in
# sibling modules; shared imports/constants/helpers stay in core.py.
from .core import (
    _md_role, redirect, messages, transaction, MaterialReference, _planner, render,
    _any_role, _make_xlsx_response, _style_xlsx_header, _add_example_rows,
    _finalize_xlsx
)
from .ref_materials_parse import parse_ref_rows, resolve_ref_columns


_REF_MATERIALS_URL = "ui:planner_ref_materials"


@_md_role
def planner_ref_materials_import(request):
    if request.method != "POST":
        return redirect(_REF_MATERIALS_URL)
    f = request.FILES.get("file")
    error = _ref_upload_error(f)
    if error:
        messages.error(request, error)
        return redirect(_REF_MATERIALS_URL)
    try:
        _import_ref_workbook(request, f)
    except Exception as exc:
        messages.error(request, f"Błąd importu: {exc}")
    return redirect(_REF_MATERIALS_URL)


def _ref_upload_error(f):
    if not f:
        return "Nie wybrano pliku."
    if f.size > 20 * 1024 * 1024:
        return "Plik zbyt duży (max 20 MB)."
    return None


def _import_ref_workbook(request, f):
    """Odczyt XLSX → parsowanie → zapis; wyjątki obsługuje widok („Błąd importu”)."""
    import openpyxl
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    if not rows:
        messages.error(request, "Plik jest pusty.")
        return
    refs, skipped = parse_ref_rows(rows[1:], resolve_ref_columns(rows[0]))
    created, updated = _save_ref_materials(refs)
    messages.success(request, f"Import zakończony: {created} nowych, {updated} zaktualizowanych, {skipped} pominiętych wierszy.")
    from ..models import ImportRun
    ImportRun.record("ref_materials", rows=created + updated,
                     label=getattr(f, "name", ""), user=request.user)


def _save_ref_materials(refs):
    """Upsert po kodzie — wszystko albo nic: błąd w połowie nie zostawia częściowego importu."""
    created = updated = 0
    with transaction.atomic():
        for code, data in refs.items():
            _obj, is_new = MaterialReference.objects.update_or_create(code=code, defaults=data)
            if is_new:
                created += 1
            else:
                updated += 1
    return created, updated


def _hdr_index(header, *needles):
    """Indeks pierwszej kolumny, której nagłówek zawiera którykolwiek z `needles`
    (case-insensitive). None, gdy brak — import fixów toleruje różne układy eksportu SAP."""
    low = [(str(h or "")).strip().lower() for h in header]
    for n in needles:
        for i, h in enumerate(low):
            if n in h:
                return i
    return None


def _num(v):
    try:
        return float(str(v).replace(" ", "").replace(",", "."))
    except (TypeError, ValueError):
        return 0.0


@_md_role
def excel_import_fix(request):
    """Import fixów (stałych lokalizacji pickingowych) z eksportu SAP: Miejsce składowania,
    Produkt, Typ magazynu, Ilość minimalna, Maksym. ilość, JM. Upsert po (lokalizacja, REF).
    Zasila „Zapas / min" na karcie produktu (MatInfo)."""
    from ..models import FixLocation
    if request.method != "POST":
        return redirect("ui:data_center")
    f = request.FILES.get("file")
    if not f:
        messages.error(request, "Nie wybrano pliku.")
        return redirect("ui:data_center")
    if f.size > 20 * 1024 * 1024:
        messages.error(request, "Plik zbyt duży (max 20 MB).")
        return redirect("ui:data_center")
    try:
        import openpyxl
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
        if not rows:
            messages.error(request, "Plik jest pusty.")
            return redirect("ui:data_center")
        h = rows[0]
        i_loc = _hdr_index(h, "miejsce składow", "miejsce", "lokaliz", "lgpla")
        i_ref = _hdr_index(h, "produkt", "materiał", "material", "ref")
        i_wt = _hdr_index(h, "typ magazynu", "lgnum", "typ")
        i_max = _hdr_index(h, "maksym", "il.maks", "il maks", "max")
        i_min = _hdr_index(h, "ilość minim", "minimalna", "il. min", "il min", "min")
        i_uom = _hdr_index(h, "jm - il. min", "jm dla il", "jm", "meins")
        i_chg = _hdr_index(h, "data zmiany")
        if i_loc is None or i_ref is None or i_min is None:
            messages.error(request, "Brak wymaganych kolumn (Miejsce składowania / Produkt / Ilość minimalna).")
            return redirect("ui:data_center")
        created = updated = skipped = 0

        def _cell(row, idx):
            # Wiersze z przyciętych eksportów SAP bywają KRÓTSZE niż nagłówek —
            # bez tego jeden krótki wiersz (IndexError) wywalał CAŁY import.
            return row[idx] if idx is not None and idx < len(row) else None

        with transaction.atomic():
            for row in rows[1:]:
                loc = (str(_cell(row, i_loc)).strip() if _cell(row, i_loc) else "")[:40]
                ref = (str(_cell(row, i_ref)).strip() if _cell(row, i_ref) else "")[:50]
                if not loc or not ref:
                    skipped += 1
                    continue
                chg = _cell(row, i_chg)
                _obj, is_new = FixLocation.objects.update_or_create(
                    location_code=loc, ref_code=ref,
                    defaults={
                        "warehouse_type": (str(_cell(row, i_wt)).strip() if _cell(row, i_wt) else "")[:10],
                        "max_qty": _num(_cell(row, i_max)) if _cell(row, i_max) is not None else 0.0,
                        "min_qty": _num(_cell(row, i_min)) if _cell(row, i_min) is not None else 0.0,
                        "uom": (str(_cell(row, i_uom)).strip() if _cell(row, i_uom) else "")[:10],
                        "changed_at": chg.date() if hasattr(chg, "date") else None,
                    })
                created += is_new
                updated += (not is_new)
        messages.success(request, f"Import fixów: {created} nowych, {updated} zaktualizowanych, {skipped} pominiętych.")
        from ..models import ImportRun
        ImportRun.record("fix_locations", rows=created + updated,
                         label=getattr(f, "name", ""), user=request.user)
    except Exception as exc:
        messages.error(request, f"Błąd importu fixów: {exc}")
    return redirect("ui:data_center")


@_planner
def planner_excel_templates(request):
    """Central page listing all available Excel import templates."""
    return render(request, "ui/planner/excel_templates.html")

@_any_role
def warehouse_combined_template(request):
    """Download a styled xlsx template for the combined upload."""
    wb, ws, response = _make_xlsx_response("PalViz_lokalizacje_combined_wzor.xlsx")

    columns = [
        # (header, description, width, example)
        ("location_code",  "Kod lokalizacji (wymagany), np. B0-01-100A",      18, "B0-01-100A"),
        ("warehouse_type", "Typ regału/lokalizacji, np. PA, PP, HH",           14, "PA"),
        ("height_mm",      "Wysokość / prześwit [mm]",                         14, 2494),
        ("width_mm",       "Szerokość fizyczna miejsca [mm]",                  14, 800),
        ("depth_mm",       "Głębokość regału [mm]",                            14, 1100),
        ("max_weight_kg",  "Maksymalna waga [kg]",                             14, 1200),
        ("max_volume_m3",  "Maksymalna objętość [m³]",                         14, 2.5),
        ("level_heights",  'Wysokości poziomów JSON, np. {"1":2494,"2":2500}', 30, '{"1":2494,"2":2500,"3":2500}'),
        ("level_cols",     'Kolumny per poziom JSON, np. {"1":1,"2":2}',       24, '{"1":1,"2":2,"3":1}'),
    ]

    _style_xlsx_header(ws, columns, title_color="1E40AF")

    example_rows = [
        ("B0-01-100A", "PA", 2494, 800, 1100, 1200, 2.5, '{"1":2494,"2":2500,"3":2500}', '{"1":1,"2":2,"3":1}'),
        ("B0-01-100X", "PA", 2500, 800, 1100, 1200, 2.5, "", ""),
        ("B0-01-100J", "PA", 2500, 800, 1100, 1200, 2.5, "", ""),
        ("B0-01-100Y", "PA", 2500, 800, 1100, 1200, 2.5, "", ""),
        ("B0-01-200A", "PP", 2200, 900, 1100, 800,  1.8, "", ""),
        ("B0-02-100A", "PP", 2200, 900, 1100, 800,  1.8, "", ""),
        ("B0-02-100X", "PP", 2200, 900, 1100, 800,  1.8, "", ""),
        ("B0-02-100J", "PP", 2200, 900, 1100, 800,  1.8, "", ""),
    ]

    _add_example_rows(ws, columns, example_rows)

    # ── Opis sheet ────────────────────────────────────────────────────────────
    from openpyxl.styles import PatternFill, Font, Alignment

    ws_desc = wb.create_sheet("Opis")
    hdr_fill = PatternFill("solid", fgColor="1E40AF")
    hdr_font = Font(bold=True, color="FFFFFF", size=11)

    ws_desc.cell(row=1, column=1, value="Kolumna").fill = hdr_fill
    ws_desc.cell(row=1, column=1).font = hdr_font
    ws_desc.cell(row=1, column=2, value="Aliasy (alternatywne nazwy kolumn)").fill = hdr_fill
    ws_desc.cell(row=1, column=2).font = hdr_font
    ws_desc.cell(row=1, column=3, value="Opis").fill = hdr_fill
    ws_desc.cell(row=1, column=3).font = hdr_font
    ws_desc.column_dimensions["A"].width = 20
    ws_desc.column_dimensions["B"].width = 50
    ws_desc.column_dimensions["C"].width = 60

    desc_data = [
        ("location_code",  "lokalizacja, adres, miejsce, location, bin",
         "Kod lokalizacji (wymagany). Format np. B0-01-100A. Używany do wyznaczenia pozycji na mapie 3D."),
        ("warehouse_type", "typ, type, rack_type, typ_lok",
         "Kod typu regału/lokalizacji. Musi istnieć w WarehouseRackType lub zostanie utworzony."),
        ("height_mm",      "wysokość, height, clearance",
         "Prześwit / wysokość miejsca w mm. Stosowany w master data lokalizacji."),
        ("width_mm",       "szerokość, width",
         "Szerokość fizyczna miejsca w mm."),
        ("depth_mm",       "głębokość, depth",
         "Głębokość regału w mm."),
        ("max_weight_kg",  "waga_max, max_weight, waga",
         "Maksymalna waga ładunku w kg."),
        ("max_volume_m3",  "objetosc_max, max_volume, objetosc",
         "Maksymalna objętość ładunku w m³."),
        ("level_heights",  "wysokosci_poziomow",
         'Wysokości poziomów jako JSON, np. {"1":2494,"2":2500,"3":2000}. Wystarczy podać w jednym wierszu danego typu.'),
        ("level_cols",     "kolumny_poziomow",
         'Liczba kolumn per poziom jako JSON, np. {"1":1,"2":2}. Wystarczy podać w jednym wierszu danego typu.'),
    ]

    for ri, (col, aliases, desc) in enumerate(desc_data, 2):
        ws_desc.cell(row=ri, column=1, value=col).font = Font(bold=True, size=10)
        ws_desc.cell(row=ri, column=2, value=aliases).font = Font(size=10, italic=True, color="374151")
        ws_desc.cell(row=ri, column=3, value=desc).font = Font(size=10)
        ws_desc.cell(row=ri, column=3).alignment = Alignment(wrap_text=True)
        ws_desc.row_dimensions[ri].height = 30

    ws.title = "Lokalizacje"

    return _finalize_xlsx(wb, ws, response)

@_any_role
def heatmap_template_download(request):
    wb, ws, response = _make_xlsx_response("PalViz_heatmapa_wzor.xlsx")
    ws.title = "Aktywności pickerów"
    cols = [
        ("location_code", "Kod lokalizacji np. B0-01-100A", 22, "B0-01-100A"),
        ("confirmed_at",  "Data i godzina potwierdzenia RRRR-MM-DD GG:MM:SS", 28, "2025-06-01 08:15:00"),
        ("picker_name",   "Imię i nazwisko pickera / login", 24, "Jan Kowalski"),
        ("task_type",     "Typ zadania np. PICK, PUT, COUNT", 18, "PICK"),
    ]
    _style_xlsx_header(ws, cols, "0F766E")
    rows = [
        ("B0-01-100A", "2025-06-01 08:15:00", "Jan Kowalski", "PICK"),
        ("B0-02-200B", "2025-06-01 08:22:00", "Anna Nowak",   "PICK"),
        ("B0-01-100A", "2025-06-01 09:05:00", "Jan Kowalski", "PICK"),
        ("C1-03-300A", "2025-06-01 10:30:00", "Piotr Wiśniewski", "COUNT"),
        ("B0-02-200B", "2025-06-02 07:55:00", "Anna Nowak",   "PUT"),
    ]
    _add_example_rows(ws, cols, rows)
    return _finalize_xlsx(wb, ws, response)

__all__ = [
    'planner_ref_materials_import',
    'excel_import_fix',
    'planner_excel_templates',
    'warehouse_combined_template',
    'heatmap_template_download',
]
