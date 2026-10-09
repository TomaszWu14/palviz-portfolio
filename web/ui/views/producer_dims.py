"""Moduł „Wymiary producenta": import deklarowanych wymiarów kartonu od producentów
i porównanie z naszą master data. Ekran w Data Center, rola Master Data."""
import openpyxl
from django.contrib import messages
from django.core.paginator import Paginator
from django.db import transaction
from django.shortcuts import redirect, render
from django.urls import reverse
from django.views.decorators.http import require_POST

from ..models import ProducerCartonBatch, ProducerCartonDim, ImportRun
from ..producer_dims import compare_dims, ours_dims, parse_dim, supplier_from_filename
from ..roles import _master_data
from .. import product_codes

_SHEET = "dane opakowań"
# Pozycje kolumn w szablonie (0-indeks): B=REF, C=opis, D=pouch, G=box, I/J/K=carton L/W/H,
# L=qty/karton, M=gross, N=net. Dane od wiersza 4 (idx 3; 3 wiersze nagłówka).
_C = {"ref": 1, "desc": 2, "pouch": 3, "box": 6, "l": 8, "w": 9, "h": 10,
      "qty": 11, "gross": 12, "net": 13}


def _cell(row, idx):
    return row[idx] if idx < len(row) else None


def _to_int(v):
    f = parse_dim(v)
    return int(f) if f is not None else None


def _ingest_file(f, user):
    """Parsuje jeden plik dostawcy → nowa aktywna ProducerCartonBatch (+ wiersze).
    Dezaktywuje poprzednią aktywną partię tego dostawcy. Rzuca ValueError przy złym
    nagłówku (kol B wiersza 3 ≠ 'REF')."""
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = wb[_SHEET] if _SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if len(rows) < 4:
        raise ValueError("Plik nie ma danych (brak wierszy pod nagłówkiem).")
    hdr = rows[2]                                    # wiersz 3 = subnagłówek z 'REF' w kol B
    if str(_cell(hdr, _C["ref"]) or "").strip().lower() != "ref":
        raise ValueError("Nie rozpoznano szablonu — kol. B wiersza 3 nie zawiera 'REF'.")
    supplier = supplier_from_filename(getattr(f, "name", ""))
    with transaction.atomic():
        (ProducerCartonBatch.objects
         .filter(supplier=supplier, is_active=True).update(is_active=False))
        batch = ProducerCartonBatch.objects.create(
            supplier=supplier, source_filename=getattr(f, "name", "")[:255],
            uploaded_by=user, is_active=True)
        objs = []
        for row in rows[3:]:                          # dane od wiersza 4 (idx 3)
            ref = str(_cell(row, _C["ref"]) or "").strip()
            if not ref:
                continue
            product = product_codes.resolve_product_code(ref)
            objs.append(ProducerCartonDim(
                batch=batch, supplier=supplier, ref_code=ref[:100], product=product,
                carton_l=parse_dim(_cell(row, _C["l"])),
                carton_w=parse_dim(_cell(row, _C["w"])),
                carton_h=parse_dim(_cell(row, _C["h"])),
                qty_in_carton=_to_int(_cell(row, _C["qty"])),
                gross_kg=parse_dim(_cell(row, _C["gross"])),
                net_kg=parse_dim(_cell(row, _C["net"])),
                box_size_raw=str(_cell(row, _C["box"]) or "")[:120],
                pouch_size_raw=str(_cell(row, _C["pouch"]) or "")[:120],
                description=str(_cell(row, _C["desc"]) or "")[:250]))
        ProducerCartonDim.objects.bulk_create(objs)
        batch.row_count = len(objs)
        batch.save(update_fields=["row_count"])
    ImportRun.record("producer_dims", rows=len(objs), label=supplier, user=user)
    return batch


def _raise_mismatch_tasks(batch):
    """Dla realnych rozjazdów partii twórz zdedublowane Taski (kategoria
    carton_dim_mismatch, dedup po dostawca+REF). Zwraca liczbę NOWYCH tasków."""
    from ..notifications import owner_users, _raise_task
    recipients = owner_users()
    n = 0
    rows = (batch.rows.select_related("product")
            .prefetch_related("product__instructions"))
    for r in rows:
        verdict, delta = compare_dims(
            (r.carton_l, r.carton_w, r.carton_h), ours_dims(r.product))
        if verdict != "mismatch":
            continue
        created = _raise_task(
            title=f"Rozjazd wymiarów kartonu: {r.ref_code} ({batch.supplier})",
            description=(f"Producent {batch.supplier} podaje inny wymiar kartonu niż "
                        f"master data (Δ {delta} cm). Zweryfikuj REF {r.ref_code}."),
            source_ref=r.ref_code, dedup_key=f"cartondim:{batch.supplier}:{r.ref_code}",
            url=reverse("ui:producer_dims"), recipients=recipients,
            category="carton_dim_mismatch")
        if created:
            n += 1
    return n


@_master_data
@require_POST
def producer_dims_upload(request):
    """Multi-upload plików dostawców. Każdy plik = partia; po imporcie auto-Taski."""
    files = request.FILES.getlist("files") or ([request.FILES["file"]]
                                               if request.FILES.get("file") else [])
    if not files:
        messages.error(request, "Nie wybrano plików.")
        return redirect("ui:producer_dims")
    ok, tasks, errors = 0, 0, 0
    for f in files:
        if f.size > 20 * 1024 * 1024:
            messages.error(request, f"{f.name}: plik zbyt duży (max 20 MB).")
            errors += 1
            continue
        try:
            batch = _ingest_file(f, request.user)
            tasks += _raise_mismatch_tasks(batch)
            ok += 1
        except Exception as exc:                      # zły szablon/odczyt — pomiń plik
            messages.error(request, f"{f.name}: {exc}")
            errors += 1
    if ok:
        messages.success(request, f"Zaimportowano plików: {ok}. Nowych zadań: {tasks}.")
    return redirect("ui:producer_dims")


@_master_data
def producer_dims_list(request):
    """Raport: aktywne wiersze producentów vs nasza master data, z odznaką werdyktu."""
    f_supplier = (request.GET.get("supplier") or "").strip()
    f_status = (request.GET.get("status") or "").strip()
    qs = (ProducerCartonDim.objects.filter(batch__is_active=True)
          .select_related("product").prefetch_related("product__instructions")
          .order_by("supplier", "ref_code"))
    if f_supplier:
        qs = qs.filter(supplier=f_supplier)
    counts = {"ok": 0, "mismatch": 0, "no_data": 0}
    rows = []
    for r in qs:
        ours = ours_dims(r.product)
        verdict, delta = compare_dims((r.carton_l, r.carton_w, r.carton_h), ours)
        counts[verdict] += 1
        if f_status and verdict != f_status:
            continue
        rows.append({"r": r, "ours": ours, "verdict": verdict, "delta": delta})
    suppliers = (ProducerCartonDim.objects.filter(batch__is_active=True)
                 .values_list("supplier", flat=True).distinct().order_by("supplier"))
    page_obj = Paginator(rows, 100).get_page(request.GET.get("page"))
    return render(request, "ui/producer_dims/list.html", {
        "page_obj": page_obj, "counts": counts, "suppliers": suppliers,
        "f_supplier": f_supplier, "f_status": f_status})


__all__ = ["producer_dims_upload", "producer_dims_list"]
