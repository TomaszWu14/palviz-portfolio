# Auto-split from the former monolithic views.py. Feature views live in
# sibling modules; shared imports/constants/helpers stay in core.py.
from .core import (
    neutralize_workbook,
    module_required, ManualForm, render, UploadCSVForm, Product, WarehouseLocationType,
    _md_role, require_POST, _build_pallet, messages, redirect, CartonVariant,
    Dimensions, PalletizationInstruction, _recalculate_instruction, _md_or_tr,
    render_to_string, HttpResponse, _form_max_height, PalletCalculator, _eval_layouts,
    transaction, Batch, Palletization, _as_int, _build_loc_fits, _render_panel,
    _planner, logging, csv, io, get_object_or_404, _re, Count, Q, Paginator,
    _meta_from_rec, mark_safe, _svg_thumbnail, _carton_from_rec, _compute_cog,
    _optimization_hints, JsonResponse, _transport_mgr, Shipment,
    _calc_shipment_data, openpyxl, _fig_layer_2d, _fig_json, safe_json
)
from django.utils import timezone


def _calc_index_context(request, manual_form=None):
    """Kontekst strony kalkulatora (wspólny dla GET i pełnego submitu „Przelicz” bez htmx)."""
    prefill = {}
    if manual_form is None:
        # Pre-fill ManualForm from GET params (e.g., links from carton form)
        get = request.GET
        if get.get("carton_l"): prefill["carton_l"] = get["carton_l"]
        if get.get("carton_w"): prefill["carton_w"] = get["carton_w"]
        if get.get("carton_h"): prefill["carton_h"] = get["carton_h"]
        if get.get("unit_weight_kg"): prefill["unit_weight"] = get["unit_weight_kg"]
        if get.get("pcs_per_carton"): prefill["pcs_per_carton"] = get["pcs_per_carton"]
        if get.get("carton_tare"): prefill["carton_tare"] = get["carton_tare"]
        if get.get("max_height"): prefill["max_height_total"] = get["max_height"]
        manual_form = ManualForm(initial=prefill) if prefill else ManualForm()
    return {
        "manual_form": manual_form,
        "upload_form": UploadCSVForm(),
        "active_tab": "calc",
        "products": Product.objects.filter(is_active=True).order_by("code"),
        "fit_locations": WarehouseLocationType.objects.filter(is_active=True).order_by("location_class", "name"),
        "prefilled": bool(prefill),
    }


@module_required("paletyzacja")
def planner_calc_index(request):
    return render(request, "ui/planner/calc_index.html", _calc_index_context(request))


def _calc_response(request, ctx, form=None):
    """htmx (HX-Request) → sam fragment wyniku; pełny submit bez JS → cała strona z wynikiem."""
    html = render_to_string("ui/_result.html", ctx, request=request)
    if request.headers.get("HX-Request") == "true":
        return HttpResponse(html)
    page = _calc_index_context(request, manual_form=form or ManualForm())
    page["result_html"] = mark_safe(html)  # nosec B308 - wyrenderowany szablon z autoescape
    return render(request, "ui/planner/calc_index.html", page)

@_md_role
@require_POST
def planner_calc_save_instruction(request):
    """Save current calculator state as a PalletizationInstruction."""
    from django.db import transaction as _tx
    from django.db.models import Max as _Max
    product_id = request.POST.get("product_id", "").strip()
    product_code = request.POST.get("product_code", "").strip().upper()
    instr_name = request.POST.get("instr_name", "").strip()

    try:
        pallet, meta = _build_pallet(
            request.POST.get("pallet", "EU"),
            int(request.POST.get("max_height_total") or 215),
            int(request.POST.get("max_weight") or 1000),
        )
        carton_l = int(request.POST.get("carton_l") or 40)
        carton_w = int(request.POST.get("carton_w") or 30)
        carton_h = int(request.POST.get("carton_h") or 25)
        unit_weight = float(request.POST.get("unit_weight") or 0.45)
        pcs = int(request.POST.get("pcs_per_carton") or 1)
        tare = float(request.POST.get("carton_tare") or 0.0)
        demand = max(1, int(request.POST.get("demand_pcs") or 1000))
        sku = request.POST.get("sku", "").strip().upper() or product_code or "SKU"
    except (ValueError, TypeError) as exc:
        messages.error(request, f"Błędne dane: {exc}")
        return redirect("ui:planner_calc_index")

    # Validate carton and code before entering the transaction
    if not product_id:
        code = product_code or sku
        if not code:
            messages.error(request, "Podaj SKU lub kod produktu.")
            return redirect("ui:planner_calc_index")
    else:
        code = None

    try:
        CartonVariant(
            sku=sku, variant="STD",
            dims=Dimensions(l_cm=carton_l, w_cm=carton_w, h_cm=carton_h),
            unit_weight_kg=unit_weight, pieces_per_carton=pcs,
            demand_pieces=demand, carton_tare_kg=tare, allow_rotation=True,
        ).validate()
    except ValueError as exc:
        messages.error(request, f"Błędne dane kartonu: {exc}")
        return redirect("ui:planner_calc_index")

    try:
        with _tx.atomic():
            if product_id:
                # pk nienumeryczny → PostgreSQL DataError (nie DoesNotExist) → 500. Waliduj int.
                try:
                    product = Product.objects.select_for_update().get(pk=int(product_id))
                except (Product.DoesNotExist, ValueError, TypeError):
                    messages.error(request, "Nie znaleziono produktu.")
                    return redirect("ui:planner_calc_index")
            else:
                product, _ = Product.objects.get_or_create(
                    code=code,
                    defaults={"name": instr_name or code},
                )
                product = Product.objects.select_for_update().get(pk=product.pk)

            last_v = product.instructions.aggregate(m=_Max("version"))["m"] or 0
            version = last_v + 1

            instr = PalletizationInstruction.objects.create(
                product=product, version=version,
                name=instr_name,
                pallet_code=meta["pallet_code"],
                pallet_length_cm=meta["length_cm"], pallet_width_cm=meta["width_cm"],
                max_height_total_cm=meta["max_height_total_cm"],
                pallet_base_height_cm=meta["base_height_cm"],
                max_weight_kg=meta["max_weight_kg"],
                carton_l=carton_l, carton_w=carton_w, carton_h=carton_h,
                unit_weight=unit_weight, pcs_per_carton=pcs,
                carton_tare=tare, demand_pcs=demand,
            )
    except Exception as exc:
        messages.error(request, f"Błąd zapisu: {exc}")
        return redirect("ui:planner_calc_index")
    try:
        _recalculate_instruction(instr)
        messages.success(request, f"Zapisano instrukcję paletyzacji dla {product.code} v{version}.")
    except Exception as exc:
        messages.warning(request, f"Zapisano, ale błąd obliczeń: {exc}")
    return redirect("ui:planner_instruction_detail", pk=instr.pk)

@_md_or_tr
def calculate(request):
    form = ManualForm(request.POST)
    if not form.is_valid():
        return _calc_response(request, {"error": "Błędne dane formularza."}, form)
    try:
        pallet, meta = _build_pallet(form.cleaned_data["pallet"], _form_max_height(form), form.cleaned_data["max_weight"])
        carton = CartonVariant(
            sku=form.cleaned_data["sku"].strip(), variant=form.cleaned_data["variant"].strip() or "STD",
            dims=Dimensions(l_cm=int(form.cleaned_data["carton_l"]), w_cm=int(form.cleaned_data["carton_w"]), h_cm=int(form.cleaned_data["carton_h"])),
            unit_weight_kg=float(form.cleaned_data["unit_weight"]),
            pieces_per_carton=int(form.cleaned_data["pcs_per_carton"]),
            demand_pieces=int(form.cleaned_data["demand_pcs"]),
            carton_tare_kg=float(form.cleaned_data.get("carton_tare") or 0.0),
            allow_rotation=True,
        ).validate()
        res = PalletCalculator.calculate(carton, pallet)
        layouts = _eval_layouts(carton, pallet, meta, res)
        selected = layouts[0]["name"]
        with transaction.atomic():     # batch + its row are all-or-nothing (no orphan Batch)
            batch = Batch.objects.create(
                name=f"Manual {timezone.localtime():%Y-%m-%d %H:%M:%S}",
                pallet_code=meta["pallet_code"], pallet_length_cm=meta["length_cm"],
                pallet_width_cm=meta["width_cm"], max_height_total_cm=meta["max_height_total_cm"],
                pallet_base_height_cm=meta["base_height_cm"], max_weight_kg=meta["max_weight_kg"])
            rec = Palletization.objects.create(
                batch=batch, sku=carton.sku, variant=carton.variant,
                carton_l=carton.dims.l_cm, carton_w=carton.dims.w_cm, carton_h=carton.dims.h_cm,
                unit_weight=carton.unit_weight_kg, pcs_per_carton=carton.pieces_per_carton,
                demand_pcs=carton.demand_pieces, carton_tare=carton.carton_tare_kg,
                layouts=layouts, selected_layout=selected)
        pid = rec.id
        locations = list(WarehouseLocationType.objects.filter(is_active=True).order_by("name"))
        # Inline "przymierz do lokalizacji" — fit + 3D/front visualisation per ticked location
        fit_ids = [_as_int(x, 0) for x in request.POST.getlist("fit_loc")]
        loc_fits = _build_loc_fits(carton.dims.l_cm, carton.dims.w_cm, carton.dims.h_cm,
                                   carton.unit_weight_kg, carton.pieces_per_carton, carton.carton_tare_kg, fit_ids)
        panel = _render_panel(meta, carton, layouts, selected, int(form.cleaned_data["render_layers"]),
                              pid=pid, locations=locations, loc_fits=loc_fits)
        return _calc_response(request, panel, form)
    except Exception as exc:
        return _calc_response(request, {"error": f"Błąd obliczania: {exc}"}, form)

@_planner
def whatif_calculate(request):
    try:
        pallet, meta = _build_pallet(
            request.GET.get("pallet", "EU"),
            int(request.GET.get("max_height_total", 215)),
            int(request.GET.get("max_weight", 1000)),
        )
        carton = CartonVariant(
            sku=request.GET.get("sku", "SKU").strip() or "SKU",
            variant=request.GET.get("variant", "STD").strip() or "STD",
            dims=Dimensions(l_cm=max(1, int(request.GET.get("carton_l", 40))),
                            w_cm=max(1, int(request.GET.get("carton_w", 30))),
                            h_cm=max(1, int(request.GET.get("carton_h", 25)))),
            unit_weight_kg=float(request.GET.get("unit_weight", 10)),
            pieces_per_carton=max(1, int(request.GET.get("pcs_per_carton", 1))),
            demand_pieces=max(1, int(request.GET.get("demand_pcs", 100))),
            carton_tare_kg=float(request.GET.get("carton_tare", 0)),
            allow_rotation=True,
        ).validate()
        res = PalletCalculator.calculate(carton, pallet)
        layouts = _eval_layouts(carton, pallet, meta, res)
        panel = _render_panel(meta, carton, layouts, layouts[0]["name"], int(request.GET.get("render_layers", 3)))
        return HttpResponse(render_to_string("ui/_result.html", panel, request=request))
    except Exception as e:
        logging.getLogger(__name__).warning("whatif_calculate failed: %s", e)
        from django.utils.html import escape
        return HttpResponse(f'<div class="card"><p style="color:#dc2626">Błąd: {escape(str(e))}</p></div>')

@_md_or_tr
def upload_csv(request):
    if request.method != "POST":
        return redirect("ui:planner_calc_index")
    form = UploadCSVForm(request.POST, request.FILES)
    if not form.is_valid():
        return render(request, "ui/planner/calc_index.html", {"manual_form": ManualForm(), "upload_form": form, "error": "Błędny upload.", "products": Product.objects.filter(is_active=True).order_by("code")})
    pallet, meta = _build_pallet(form.cleaned_data["pallet"], _form_max_height(form), form.cleaned_data["max_weight"])
    f = request.FILES["file"]
    if f.size > 10 * 1024 * 1024:
        return render(request, "ui/planner/calc_index.html", {"manual_form": ManualForm(), "upload_form": form, "error": "Plik zbyt duży (max 10 MB).", "products": Product.objects.filter(is_active=True).order_by("code")})
    content = f.read().decode("utf-8", errors="ignore")
    # Detect the delimiter — European Excel exports use ';' which otherwise yields a
    # single combined column and a false "missing columns" error.
    _sample = content[:2048]
    try:
        _delim = csv.Sniffer().sniff(_sample, delimiters=",;\t").delimiter
    except csv.Error:
        _delim = ";" if _sample.count(";") > _sample.count(",") else ","
    reader = csv.DictReader(io.StringIO(content), delimiter=_delim)
    required = {"SKU", "WARIANT", "L", "W", "H", "WAGA_SZT", "SZT_W_KARTONIE", "ILOSC_SZT"}
    missing = required - {c.strip().upper() for c in (reader.fieldnames or [])}
    if missing:
        return render(request, "ui/planner/calc_index.html", {"manual_form": ManualForm(), "upload_form": form, "error": f"Brak kolumn: {sorted(missing)}", "products": Product.objects.filter(is_active=True).order_by("code")})
    batch = Batch.objects.create(name=f"CSV {f.name}"[:250], pallet_code=meta["pallet_code"],
        pallet_length_cm=meta["length_cm"], pallet_width_cm=meta["width_cm"],
        max_height_total_cm=meta["max_height_total_cm"], pallet_base_height_cm=meta["base_height_cm"],
        max_weight_kg=meta["max_weight_kg"])
    errors = []
    for row_num, row in enumerate(reader, start=2):
        try:
            def ff(x): return float(str(x).strip().replace(",", "."))
            def fi(x):
                # Jak parse_int z palletizer.io.parsing: „12,7" to błąd danych, nie 12 —
                # ciche obcinanie dawało inne kartony niż loader CLI z tego samego pliku.
                v = float(str(x).strip().replace(",", "."))
                if v != int(v):
                    raise ValueError(f"wartość całkowita wymagana, jest: {x}")
                return int(v)
            tare_key = next((k for k in row if k.upper() == "KARTON_TARE"), None)
            carton = CartonVariant(
                sku=str(row.get("SKU", "")).strip(), variant=str(row.get("WARIANT", "STD")).strip() or "STD",
                dims=Dimensions(l_cm=fi(row["L"]), w_cm=fi(row["W"]), h_cm=fi(row["H"])),
                unit_weight_kg=ff(row["WAGA_SZT"]), pieces_per_carton=fi(row["SZT_W_KARTONIE"]),
                demand_pieces=fi(row["ILOSC_SZT"]), carton_tare_kg=ff(row.get(tare_key, 0) or 0) if tare_key else 0.0,
                allow_rotation=True).validate()
            res = PalletCalculator.calculate(carton, pallet)
            layouts = _eval_layouts(carton, pallet, meta, res)
            Palletization.objects.create(batch=batch, sku=carton.sku, variant=carton.variant,
                carton_l=carton.dims.l_cm, carton_w=carton.dims.w_cm, carton_h=carton.dims.h_cm,
                unit_weight=carton.unit_weight_kg, pcs_per_carton=carton.pieces_per_carton,
                demand_pcs=carton.demand_pieces, carton_tare=carton.carton_tare_kg,
                layouts=layouts, selected_layout=layouts[0]["name"])
        except Exception as e:
            errors.append(f"Wiersz {row_num}: {e}")
    if errors:
        messages.warning(request, "Batch zapisany z błędami: " + "; ".join(errors[:3]))
    return redirect("ui:batch_detail", batch_id=batch.id)

@_planner
def export_excel(request, pid: int):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return HttpResponse("Brak openpyxl.", status=500)
    rec = get_object_or_404(Palletization, id=pid)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{rec.sku}_{rec.variant}"[:31]
    hf = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="1A56DB")
    bfill = PatternFill("solid", fgColor="D1FAE5")
    headers = ["Layout","Kart/warstwa","Warstwy","Kart/paleta","Pełne palety","Reszta","Wys.cargo","Wys.total","Waga/pal","Area%","Cube%","CoG X","CoG Y","ΔCoG X","ΔCoG Y","Najlepszy"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = hf; cell.fill = hfill; cell.alignment = Alignment(horizontal="center")
    for lay in rec.layouts:
        cog = lay.get("cog", {})
        # .get throughout — a legacy or OR-Tools-shaped layout row may lack some keys, and a
        # direct lay["..."] would 500 the whole export (the cog access below already uses .get).
        row = [lay.get("name",""),lay.get("cartons_per_layer",""),lay.get("layers_used",""),
               lay.get("cartons_per_pallet",""),lay.get("pallets_full",""),lay.get("remainder_cartons",""),
               lay.get("cargo_height_used_cm",""),lay.get("total_height_used_cm",""),
               lay.get("weight_per_pallet_kg",""),lay.get("area_used_pct",""),lay.get("cube_used_pct",""),
               cog.get("x",""),cog.get("y",""),cog.get("offset_x",""),cog.get("offset_y",""),
               "TAK" if lay.get("is_best") else ""]
        ws.append(row)
        if lay.get("is_best"):
            for cell in ws[ws.max_row]: cell.fill = bfill
    output = io.BytesIO()
    neutralize_workbook(wb)  # SEC-007
    wb.save(output); output.seek(0)
    _safe_sku = _re.sub(r'[^\w\-.]', '_', rec.sku)
    _safe_var = _re.sub(r'[^\w\-.]', '_', rec.variant)
    fname = f"paletyzacja_{_safe_sku}_{_safe_var}.xlsx"
    resp = HttpResponse(output.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f"attachment; filename=\"{fname}\""
    return resp

@require_POST
@_md_or_tr
def delete_batch(request, batch_id: int):
    get_object_or_404(Batch, id=batch_id).delete()
    messages.success(request, "Batch usunięty.")
    return redirect("ui:saved_list")

@_planner
def saved_list(request):
    q = request.GET.get("q", "").strip()
    qs = Batch.objects.annotate(items_cnt=Count("items", distinct=True)).order_by("-created_at")
    if q:
        qs = qs.filter(Q(name__icontains=q) | Q(pallet_code__icontains=q) | Q(items__sku__icontains=q)).distinct()
    page = Paginator(qs, 20).get_page(request.GET.get("page", 1))
    return render(request, "ui/saved_list.html", {"batches": page, "q": q})

@_planner
def batch_detail(request, batch_id: int):
    batch = get_object_or_404(Batch, id=batch_id)
    return render(request, "ui/batch_detail.html", {"batch": batch, "items": batch.items.order_by("sku", "variant")})

@_planner
def pallet_detail(request, pid: int):
    rec = get_object_or_404(Palletization, id=pid)
    meta = _meta_from_rec(rec)
    selected = rec.selected_layout or (rec.layouts[0]["name"] if rec.layouts else "")
    return render(request, "ui/pallet_detail.html", {
        "rec": rec, "meta": meta, "selected": selected,
        "top10": rec.layouts[:15],
        "thumbs": [(l["name"], mark_safe(_svg_thumbnail(meta, l))) for l in rec.layouts[:15]],
    })

@_planner
def pallet_panel(request, pid: int):
    rec = get_object_or_404(Palletization, id=pid)
    layout_name = request.GET.get("layout", "") or rec.selected_layout
    render_layers = max(1, min(_as_int(request.GET.get("render_layers"), 3), 20))
    layer_idx = max(1, _as_int(request.GET.get("layer_idx"), 1))
    meta = _meta_from_rec(rec)
    carton = _carton_from_rec(rec)
    _layouts = rec.layouts or []
    layout = next((l for l in _layouts if l.get("name") == layout_name), None) or (_layouts[0] if _layouts else None)
    if not layout:
        return HttpResponse("<div class='card'>Brak layoutów.</div>")
    # Legacy/OR-Tools-shaped layout rows may lack 'placements' — guard like export_excel does.
    _placements = layout.get("placements") or []
    if "cog" not in layout:
        layout["cog"] = _compute_cog(_placements, meta)
    if "hints" not in layout:
        layout["hints"] = _optimization_hints(layout, rec.carton_h, meta["cargo_max_height_cm"], meta["max_weight_kg"])
    layer_idx = max(1, min(layer_idx, layout.get("layers_used", 1)))
    locations = list(WarehouseLocationType.objects.filter(is_active=True).order_by("name"))
    # Keep the inline location fits alive when switching the pallet variant (the checkboxes
    # are hx-included, so loc_fits don't reset and re-rendering happens live).
    fit_ids = [_as_int(x, 0) for x in request.GET.getlist("fit_loc")]
    loc_fits = _build_loc_fits(carton.dims.l_cm, carton.dims.w_cm, carton.dims.h_cm,
                               carton.unit_weight_kg, carton.pieces_per_carton, carton.carton_tare_kg, fit_ids)
    payload = _render_panel(meta, carton, rec.layouts, layout_name, render_layers,
                            pid=pid, locations=locations, loc_fits=loc_fits)
    payload["layer_idx"] = layer_idx
    payload["layout"] = layout
    payload["placements_json"] = safe_json(_placements)
    payload["carton_dims_json"] = safe_json({"h": carton.dims.h_cm, "l": carton.dims.l_cm, "w": carton.dims.w_cm})
    return HttpResponse(render_to_string("ui/_panel.html", payload, request=request))

@_md_role      # B-003: przeliczenie = zapis instrukcji (instructions_write: Admin/MD)
@require_POST
def planner_instruction_recalculate_async(request, pk: int):
    """Trigger async recalculation of one instruction; return task_id."""
    from ..tasks import recalculate_instruction_task
    get_object_or_404(PalletizationInstruction, pk=pk)
    task = recalculate_instruction_task.delay(pk)
    return JsonResponse({"task_id": task.id, "instr_pk": pk})

@_md_role      # B-003: przeliczenie = zapis instrukcji (instructions_write: Admin/MD)
@require_POST
def recalculate_all(request):
    """Trigger async batch recalculation of all active instructions."""
    from ..tasks import recalculate_all_task
    task = recalculate_all_task.delay()
    return JsonResponse({"task_id": task.id})

@_transport_mgr
def planner_shipment_calc_api(request, pk):
    """AJAX: return calculation JSON for a saved shipment."""
    shipment = get_object_or_404(Shipment, pk=pk)
    calc = _calc_shipment_data(shipment)
    return JsonResponse({
        "total_vol_m3": calc["total_vol_m3"],
        "total_weight_kg": calc["total_weight_kg"],
        "total_cartons": calc["total_cartons"],
        "scenarios": calc["scenarios"],
        "errors": calc["errors"],
    })

@_transport_mgr
def planner_shipment_export_excel(request, pk):
    """Export shipment calculation to Excel."""
    shipment = get_object_or_404(Shipment, pk=pk)
    calc = _calc_shipment_data(shipment)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Przesyłka"
    ws.append(["Przesyłka", shipment.name])
    ws.append(["Kraj", shipment.destination_country])
    ws.append(["Miasto", shipment.destination_city])
    ws.append([])
    ws.append(["Kod", "Nazwa", "Ilość", "Jedn.", "Kartony", "Obj. [m³]", "Waga [kg]"])
    for lc in calc["lines"]:
        ws.append([
            lc["product"].code,
            lc["product"].name,
            lc["line"].quantity,
            lc["line"].get_unit_display(),
            lc["n_cartons"],
            lc["volume_m3"],
            lc["weight_kg"],
        ])
    ws.append([])
    ws.append(["RAZEM", "", "", "", calc["total_cartons"], calc["total_vol_m3"], calc["total_weight_kg"]])
    ws.append([])
    ws.append(["Scenariusz", "Palet", "LDM", "Wypełnienie"])
    for sc in calc["scenarios"]:
        ws.append([sc["label"], sc["n_pallets"], sc["ldm"], f"{sc['fill_pct']}%"])

    buf = io.BytesIO()
    neutralize_workbook(wb)  # SEC-007
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(buf.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="przesylka_{shipment.pk}.xlsx"'
    return resp

@_planner
def calc_optimal_layer(request, pid: int):
    """On-demand: prove the optimal cartons-per-layer for a saved pallet via OR-Tools
    CP-SAT and compare it to the best heuristic. Opt-in (a button in the result panel) so
    the solver's time budget never slows the main calc / bulk-CSV flows."""
    from palletizer.services.ortools_layer import optimal_uniform_layer

    rec = get_object_or_404(Palletization, pk=pid)
    meta = _meta_from_rec(rec)
    L, W = int(meta["length_cm"]), int(meta["width_cm"])
    cl, cw = int(rec.carton_l), int(rec.carton_w)
    placements = optimal_uniform_layer(L, W, cl, cw, allow_rotation=True, time_limit_s=4.0)
    if not placements:
        return HttpResponse(render_to_string("ui/_optimal_layer.html", {"unavailable": True}, request=request))

    layout = {
        "placements": [{"x": x, "y": y, "dx": w, "dy": h, "rotated": (w != cl)}
                       for (x, y, w, h) in placements],
        "layers_used": 1, "cog": {},
    }
    fig = _fig_layer_2d(meta, layout, rec.sku, 1)
    best = max((int(l.get("cartons_per_layer", 0)) for l in (rec.layouts or [])), default=0)
    opt = len(placements)
    area_pct = round(opt * cl * cw / (L * W) * 100.0, 1) if L and W else 0.0
    ctx = {
        "fig_json": _fig_json(fig),
        "optimal": opt, "best_heuristic": best, "delta": opt - best, "area_pct": area_pct,
    }
    return HttpResponse(render_to_string("ui/_optimal_layer.html", ctx, request=request))

__all__ = [
    'planner_calc_index',
    'planner_calc_save_instruction',
    'calculate',
    'whatif_calculate',
    'upload_csv',
    'export_excel',
    'delete_batch',
    'saved_list',
    'batch_detail',
    'pallet_detail',
    'pallet_panel',
    'calc_optimal_layer',
    'planner_instruction_recalculate_async',
    'recalculate_all',
    'planner_shipment_calc_api',
    'planner_shipment_export_excel',
]
