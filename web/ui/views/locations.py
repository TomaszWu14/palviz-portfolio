# Auto-split from the former monolithic views.py. Feature views live in
# sibling modules; shared imports/constants/helpers stay in core.py.
from .core import (
    _md_role, redirect, messages, _read_xlsx_as_dicts, date,
    WarehouseLocationMasterBatch, WarehouseLocationMaster, _planner,
    _make_xlsx_response, _style_xlsx_header, _add_example_rows, _finalize_xlsx,
    WarehouseLocationType, render, get_object_or_404, _fig_location_3d, _fig_json,
    _fig_location_2d_front
)
import math

# WarehouseLocationMaster.height_mm to IntegerField(default=0) → int4 w Postgresie.
_INT4_MAX = 2_147_483_647
_MAX_ERROR_DETAILS = 10


def _parse_capacity_mm(raw):
    """Pojemność [mm] z komórki xlsx → (wartość, błąd).

    Puste → 0 (domyślna modelu). Tekst nieparsowalny, inf/NaN, wartość ujemna lub
    poza zakresem int4 → (None, opis błędu) — wiersz ma zostać odrzucony. Wcześniej
    taka wartość stawała się None w kolumnie NOT NULL (IntegrityError) albo rzucała
    OverflowError (int(inf)) i przerywała CAŁY import."""
    s = (raw or "").strip()
    if not s:
        return 0, None
    try:
        v = float(s.replace(",", ".").replace(" ", "").replace(" ", ""))
    except ValueError:
        return None, "nie jest liczbą"
    if not math.isfinite(v):
        return None, "wartość nieskończona / NaN"
    if v < 0 or v > _INT4_MAX:
        return None, f"poza zakresem 0–{_INT4_MAX:,}".replace(",", " ")
    return int(v), None


@_md_role
def excel_import_locations(request):
    """Import warehouse location master data from xlsx template (creates a new batch)."""
    if request.method != "POST":
        return redirect("ui:planner_excel_templates")

    f = request.FILES.get("file")
    if not f:
        messages.error(request, "Nie wybrano pliku.")
        return redirect("ui:planner_excel_templates")
    if f.size > 20 * 1024 * 1024:
        messages.error(request, "Plik zbyt duży (max 20 MB).")
        return redirect("ui:planner_excel_templates")

    try:
        from django.db import transaction
        rows = _read_xlsx_as_dicts(f, row_numbers=True)
        batch_name = f"Import Excel {date.today()}"
        errors = 0
        error_details = []
        locations_data = []
        for row in rows:
            loc = row.get("location_code", "").strip()[:50]   # CharField(50) — Postgres 500 przy dłuższym
            if not loc:
                errors += 1
                continue
            cap_mm, cap_err = _parse_capacity_mm(row.get("capacity_mm"))
            if cap_err:
                errors += 1
                error_details.append(
                    f"wiersz {row['_row']} ({loc}), kolumna capacity_mm: "
                    f"„{row.get('capacity_mm', '')[:40]}” — {cap_err}")
                continue
            locations_data.append(dict(
                location_code=loc,
                warehouse_type=row.get("warehouse_type", "").strip()[:50],
                height_mm=cap_mm,
            ))

        with transaction.atomic():
            # Deactivate all previous batches so the new one becomes the active source
            WarehouseLocationMasterBatch.objects.filter(is_active=True).update(is_active=False)
            batch = WarehouseLocationMasterBatch.objects.create(
                name=batch_name, location_count=len(locations_data)
            )
            WarehouseLocationMaster.objects.bulk_create([
                WarehouseLocationMaster(
                    batch=batch,
                    location_code=d["location_code"],
                    warehouse_type=d["warehouse_type"],
                    height_mm=d["height_mm"],
                )
                for d in locations_data
            ])

        msg = f"Import lokalizacji: {len(locations_data)} wpisów, {errors} błędów. Batch: '{batch_name}'."
        if errors:
            messages.warning(request, msg)
            if error_details:
                more = len(error_details) - _MAX_ERROR_DETAILS
                tail = f" (+{more} kolejnych)" if more > 0 else ""
                messages.warning(
                    request,
                    "Odrzucone wiersze: " + "; ".join(error_details[:_MAX_ERROR_DETAILS]) + tail + ".")
        else:
            messages.success(request, msg)
    except Exception as e:
        messages.error(request, f"Błąd wczytywania pliku: {e}")

    return redirect("ui:planner_locations")

@_planner
def excel_template_locations(request):
    wb, ws, response = _make_xlsx_response("PalViz_lokalizacje_magazynowe_wzor.xlsx")
    ws.title = "Lokalizacje"
    cols = [
        ("location_code",  "Kod lokalizacji — format Strefa-Regał-Bok+Poziom", 22, "B0-01-300A"),
        ("warehouse_type", "Typ magazynu / strefa (opcjonalnie)",               18, "B0"),
        ("capacity_mm",    "Pojemność (maks. wys. towaru) [mm] (opcj.)",        18, 2100),
        ("notes",          "Uwagi (opcjonalnie)",                               24, ""),
    ]
    _style_xlsx_header(ws, cols, "0F766E")
    rows = [
        ["B0-01-300A","B0",2100,""],["B0-01-300B","B0",2100,""],["B0-01-300C","B0",2100,""],
        ["B0-01-300X","B0",1800,"Poz.2 slot0"],["B0-01-300J","B0",1800,"Poz.2 slot1"],["B0-01-300K","B0",1800,"Poz.2 slot2"],
        ["B0-02-100A","B0",2100,""],["B0-02-100Y","B0",1800,""],["B0-02-100L","B0",1800,"Poz.3 slot1"],
        ["C1-01-200A","C1",2200,""],["C1-01-200Z","C1",1500,""],["C1-01-200N","C1",1500,"Poz.4 slot1"],
        ["A0-01-500A","A0",1800,"Strefa wysokiego składowania"],
    ]
    _add_example_rows(ws, cols, rows)
    # Second sheet: format explanation
    ws2 = wb.create_sheet("Format kodu")
    from openpyxl.styles import Font, PatternFill
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 60
    headers = [
        ("Element kodu", "Opis i przykład"),
        ("B0",  "Strefa (litera+cyfra, np. A0, B1, C2)"),
        ("01",  "Numer regału w strefie (2 cyfry)"),
        ("300", "Numer boku/kolumny (3 cyfry)"),
        ("Sufiks", "Poziom i slot — tabela poniżej"),
    ]
    for r, (col1, col2) in enumerate(headers, 1):
        c1 = ws2.cell(row=r, column=1, value=col1)
        c2 = ws2.cell(row=r, column=2, value=col2)
        if r == 1:
            c1.font = Font(bold=True, color="FFFFFF")
            c2.font = Font(bold=True, color="FFFFFF")
            c1.fill = PatternFill("solid", fgColor="0F766E")
            c2.fill = PatternFill("solid", fgColor="0F766E")
        else:
            c1.font = Font(bold=True, size=11)
    # Suffix table
    ws2.cell(row=7, column=1, value="Tabela sufiksów:").font = Font(bold=True)
    suffix_rows = [
        ("A", "Poziom 1, slot 0 (lewa kolumna lub jedyna)"),
        ("B", "Poziom 1, slot 1 (środkowa kolumna)"),
        ("C", "Poziom 1, slot 2 (prawa kolumna)"),
        ("X", "Poziom 2, slot 0"),
        ("J", "Poziom 2, slot 1 — NOWY"),
        ("K", "Poziom 2, slot 2 — NOWY"),
        ("Y", "Poziom 3, slot 0"),
        ("L", "Poziom 3, slot 1 — NOWY"),
        ("M", "Poziom 3, slot 2 — NOWY"),
        ("Z", "Poziom 4, slot 0"),
        ("N", "Poziom 4, slot 1 — NOWY"),
        ("O", "Poziom 4, slot 2 — NOWY"),
    ]
    for i, (suffix, desc) in enumerate(suffix_rows, 8):
        c1 = ws2.cell(row=i, column=1, value=suffix)
        c1.font = Font(bold=True, size=11, color="0F766E")
        ws2.cell(row=i, column=2, value=desc).font = Font(size=10)
    ws2.cell(row=21, column=1, value="Przykład pełny:").font = Font(bold=True)
    ws2.cell(row=21, column=2, value="B0-01-300J = strefa B0, regał 01, bok 300, poziom 2 slot 1 (nowy)").font = Font(size=10)
    return _finalize_xlsx(wb, ws, response)

@_planner
def planner_locations(request):
    locs = WarehouseLocationType.objects.all()
    return render(request, "ui/planner/locations.html", {"locations": locs})

@_md_role
def planner_location_form(request, pk=None):
    from ..forms import WarehouseLocationTypeForm
    instance = get_object_or_404(WarehouseLocationType, pk=pk) if pk else None
    if request.method == "POST":
        form = WarehouseLocationTypeForm(request.POST, instance=instance)
        if form.is_valid():
            form.save()
            messages.success(request, "Lokalizacja zapisana." if pk else "Lokalizacja dodana.")
            return redirect("ui:planner_locations")
    else:
        form = WarehouseLocationTypeForm(instance=instance)
    return render(request, "ui/planner/location_form.html", {"form": form, "instance": instance})

@_md_role
def planner_location_delete(request, pk):
    loc = get_object_or_404(WarehouseLocationType, pk=pk)
    if request.method == "POST":
        loc.delete()
        messages.success(request, f'Lokalizacja "{loc.name}" usunięta.')
        return redirect("ui:planner_locations")
    return render(request, "ui/planner/location_confirm_delete.html", {"loc": loc})

@_planner
def planner_location_simulate(request, pk):
    from ..forms import LocationSimulateForm
    loc = get_object_or_404(WarehouseLocationType, pk=pk)
    with_pallet = request.GET.get("with_pallet", "1") != "0"
    form = LocationSimulateForm(request.GET or None)
    fig_json = fig_2d_json = fit = None
    if form.is_valid():
        d = form.cleaned_data
        carton_from_db = d.get("carton_from_db")
        if carton_from_db:
            cl, cw, ch = carton_from_db.length_cm, carton_from_db.width_cm, carton_from_db.height_cm
            uw, ppc, tare = carton_from_db.unit_weight_kg, carton_from_db.pieces_per_carton, carton_from_db.tare_kg
        else:
            cl = d["carton_l"]; cw = d["carton_w"]; ch = d["carton_h"]
            uw = d["unit_weight"]; ppc = d["pcs_per_carton"]; tare = d.get("carton_tare") or 0.0
        # Use location's pallet flag as default if not explicitly passed
        if request.GET.get("with_pallet") is None:
            with_pallet = loc.is_pallet_location
        fig, fit = _fig_location_3d(loc, cl, cw, ch, uw, ppc, tare, with_pallet=with_pallet)
        fig_json = _fig_json(fig)
        fig_2d = _fig_location_2d_front(loc, cw, ch, fit)
        fig_2d_json = _fig_json(fig_2d)
    return render(request, "ui/planner/location_simulate.html",
                  {"loc": loc, "form": form, "fig_json": fig_json,
                   "fig_2d_json": fig_2d_json, "fit": fit,
                   "with_pallet": with_pallet})

@_planner
def planner_location_fit_all(request):
    """Compare carton fit across all active locations (with and without pallet)."""
    from ..forms import LocationSimulateForm
    form = LocationSimulateForm(request.GET or None)
    results = []
    cl = cw = ch = uw = ppc = tare = None

    if form.is_valid():
        d = form.cleaned_data
        carton_from_db = d.get("carton_from_db")
        if carton_from_db:
            cl, cw, ch = carton_from_db.length_cm, carton_from_db.width_cm, carton_from_db.height_cm
            uw, ppc, tare = carton_from_db.unit_weight_kg, carton_from_db.pieces_per_carton, carton_from_db.tare_kg
        else:
            cl = d["carton_l"]; cw = d["carton_w"]; ch = d["carton_h"]
            uw = d["unit_weight"]; ppc = d["pcs_per_carton"]; tare = d.get("carton_tare") or 0.0

        locs = WarehouseLocationType.objects.filter(is_active=True).order_by("location_class", "name")
        for loc in locs:
            row = {"loc": loc, "modes": []}
            # Primary mode = driven by location type; secondary = alternative view
            primary_pallet = loc.is_pallet_location
            for with_pallet in ([True, False] if primary_pallet else [False, True]):
                fig3d, fit = _fig_location_3d(loc, cl, cw, ch, uw, ppc, tare, with_pallet=with_pallet)
                fig2d = _fig_location_2d_front(loc, cw, ch, fit)
                row["modes"].append({
                    "with_pallet": with_pallet,
                    "fit": fit,
                    "fig3d_json": _fig_json(fig3d),
                    "fig2d_json": _fig_json(fig2d),
                })
            results.append(row)

    return render(request, "ui/planner/location_fit_all.html", {
        "form": form, "results": results,
        "cl": cl, "cw": cw, "ch": ch, "uw": uw, "ppc": ppc,
        "active_tab": "locations",
    })

__all__ = [
    'excel_import_locations',
    'excel_template_locations',
    'planner_locations',
    'planner_location_form',
    'planner_location_delete',
    'planner_location_simulate',
    'planner_location_fit_all',
]
