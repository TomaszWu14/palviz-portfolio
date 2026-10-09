"""Wzór „Dane do kontroli HU" ma nagłówki 1:1 jak realny eksport SAP (20 kolumn) i
przechodzi przez importer bez rozjazdu (round-trip)."""
import io

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import TestCase
from django.urls import reverse

from ui.models import HandlingUnit
from huctl.views.hu import _import_hu_rows

EXPECTED = [
    "Jednostka obsługi", "Produkt", "Krótki opis produktu", "Partia dostawcy",
    "Termin ważności", "Ilość PJM", "Podst. jedn. miary", "Jednostka alternat.",
    "Ilość AJM", "Zapas_HU.Miejsce składowania", "Dokument",
    "Nagłówki.Odbiorca materiałów", "Nagłówki.Utworz. dnia", "Nagłówki.Autor",
    "Nagłówki.Klucz kraju/regionu", "Nagłówki.Status pobrania", "Potwierdzone przez",
    "Nagłówki.Nazwa odbiorcy 1", "Nagłówki.Utworzono o godz.", "Lokalizacje.Typ magazynu",
]


class ControlTemplateSapTests(TestCase):
    def setUp(self):
        u = get_user_model().objects.create_user(username="u", password="x")
        u.groups.add(Group.objects.get_or_create(name="Podgląd")[0])  # _planner wymaga roli
        self.client.force_login(u)

    def _workbook(self):
        r = self.client.get(reverse("ui:excel_template_control_data"))
        self.assertEqual(r.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(r.content))

    def test_headers_match_sap_20_columns(self):
        ws = self._workbook().worksheets[0]
        hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        self.assertEqual(hdr, EXPECTED)

    def test_template_rows_round_trip_through_importer(self):
        ws = self._workbook().worksheets[0]
        header = [str(ws.cell(1, c).value).lower() for c in range(1, ws.max_column + 1)]
        rows = [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
                for r in range(2, ws.max_row + 1)]
        self.assertTrue(rows, "wzór ma przykładowe wiersze")
        ok, info = _import_hu_rows(header, rows)
        self.assertTrue(ok, info)
        hu = HandlingUnit.objects.get(code="11582503")
        self.assertEqual(hu.warehouse_type, "92EX")     # Typ magazynu (gating)
        self.assertEqual(hu.picker, "JKOWAL")            # Potwierdzone przez
        self.assertFalse(hu.shipment.is_stock)          # Dokument → dostawa
