"""Ekrany „Wykryj z EWM” (podgląd → zapis) i „Zgodność z EWM” (+ XLSX), z rolami."""
import io

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import TestCase
from django.urls import reverse

from wh3d.models import BayTemplate, LocationOverride, WarehouseLocationMasterBatch
from wh3d.tests.test_ewm_service import make_model_and_master


class EwmViewsTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_superuser(username="adm", password="x")
        self.viewer = User.objects.create_user(username="podglad", password="x")
        self.viewer.groups.add(Group.objects.get_or_create(name="Podgląd")[0])
        self.wm, self.batch = make_model_and_master(extra_codes=["B0-60-100A"])

    def test_detect_preview_lists_templates_and_rows_without_saving(self):
        self.client.force_login(self.viewer)
        r = self.client.get(reverse("ui:warehouse_model_detect", args=[self.wm.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "3 pal. · B C½ D½ X Y Z · 0052/0010")
        self.assertContains(r, "10-47,50")
        self.assertNotContains(r, reverse("ui:warehouse_model_detect_save", args=[self.wm.pk]))  # brak roli MD
        self.assertEqual(BayTemplate.objects.count(), 0)

    def test_viewer_cannot_save(self):
        self.client.force_login(self.viewer)
        r = self.client.post(reverse("ui:warehouse_model_detect_save", args=[self.wm.pk]))
        self.assertEqual(r.status_code, 403)
        self.assertEqual(BayTemplate.objects.count(), 0)

    def test_save_then_compliance_shows_ok_and_no_row(self):
        self.client.force_login(self.admin)
        r = self.client.post(reverse("ui:warehouse_model_detect_save", args=[self.wm.pk]), follow=True)
        self.assertRedirects(r, reverse("ui:warehouse_model_compliance", args=[self.wm.pk]))
        self.assertEqual(BayTemplate.objects.count(), 18)
        self.assertContains(r, "Zgodne")
        self.assertContains(r, "Brak rzędu na planie")              # przejście 60 tylko w EWM
        self.assertContains(r, "6 / 7")                              # zgodne przejścia / wszystkie

    def test_compliance_xlsx(self):
        self.client.force_login(self.viewer)
        r = self.client.get(reverse("ui:warehouse_model_compliance", args=[self.wm.pk]) + "?format=xlsx")
        self.assertEqual(r["Content-Type"],
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Zgodność", "Rozbieżności"])
        self.assertEqual(wb["Zgodność"]["A1"].value, "Strefa")

    def test_compliance_xlsx_escapes_formula_injection(self):
        self.client.force_login(self.admin)
        self.client.post(reverse("ui:warehouse_model_detect_save", args=[self.wm.pk]))
        rack07 = self.wm.racks.get(rack_id="07")
        LocationOverride.objects.create(rack=rack07, bay=10, position=0, letter="A", half=0,
                                        action="rename", value="=1+1")
        r = self.client.get(reverse("ui:warehouse_model_compliance", args=[self.wm.pk]) + "?format=xlsx")
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        diff = wb["Rozbieżności"]
        values = [cell.value for row in diff.iter_rows() for cell in row]
        self.assertIn("'=1+1", values)

    def test_no_active_master_disables_detect(self):
        WarehouseLocationMasterBatch.objects.update(is_active=False)
        self.client.force_login(self.admin)
        r = self.client.get(reverse("ui:warehouse_model_detect", args=[self.wm.pk]))
        self.assertContains(r, "Brak aktywnego mastera lokalizacji")
        r = self.client.post(reverse("ui:warehouse_model_detect_save", args=[self.wm.pk]), follow=True)
        self.assertContains(r, "Brak aktywnego mastera lokalizacji")
        self.assertEqual(BayTemplate.objects.count(), 0)

    def test_model_view_links(self):
        self.client.force_login(self.admin)
        r = self.client.get(reverse("ui:warehouse_model_view", args=[self.wm.pk]))
        self.assertContains(r, reverse("ui:warehouse_model_detect", args=[self.wm.pk]))
        self.assertContains(r, reverse("ui:warehouse_model_compliance", args=[self.wm.pk]))

    def test_model_view_detect_button_disabled_without_master(self):
        WarehouseLocationMasterBatch.objects.update(is_active=False)
        self.client.force_login(self.admin)
        r = self.client.get(reverse("ui:warehouse_model_view", args=[self.wm.pk]))
        self.assertNotContains(r, reverse("ui:warehouse_model_detect", args=[self.wm.pk]))
        self.assertContains(r, "Brak aktywnego mastera lokalizacji")
