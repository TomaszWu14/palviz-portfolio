"""Jedno mapowanie liter EWM → poziom (ui.views.core.ewm_levels) i jego konsumenci.

Reguły z realnego eksportu EWM (35 780 lokalizacji regałowych, potwierdzone przez właściciela):
hala B — poz. 1: A albo półki B/C/D (jedna nad drugą), S; poz. 2: X albo G+H (X podzielone
pionowo, G niżej), T; poz. 3: Y/U; poz. 4: Z/V; poz. 5: W. Hala A (A0–A3): A–E = poziomy 1–5.
"""
from pathlib import Path

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from django.urls import reverse

from ui.management.commands.import_locations import parse_rows
from ui.models import WarehouseLayout, WarehouseLayoutCell, WarehouseLocationMaster, WarehouseRackType
from ui.roles import ALL_GROUPS
from ui.views.core.ewm_levels import (
    LetterSlot, code_slot, letter_level, letter_slot, level_height_keys, level_height_mm,
)
from ui.views.core.helpers import _parse_loc_code, _parse_location_code
from wh3d.addressing import LETTER_ORDER
from wh3d.model_from_layout import plan_racks
from wh3d.views.warehouse_map_upload_parse import parse_location

# kod → (poziom, półka, strona)
CASES = {
    # hala B — poziom 1
    "B0-08-221A": (1, 0, 0),
    "B0-39-300B": (1, 1, 0),
    "B0-39-300C": (1, 2, 0),
    "B0-39-300D": (1, 3, 0),
    "B0-07-300C-1": (1, 2, 1),
    "B0-07-300C-2": (1, 2, 2),
    "B0-07-300D-1": (1, 3, 1),
    "B0-07-300D-2": (1, 3, 2),
    "B0-76-700S": (1, 0, 0),
    # poziom 2: X albo G/H (X podzielone pionowo: G niżej, H wyżej) albo T
    "B0-08-221X": (2, 0, 0),
    "B0-55-700G": (2, 1, 0),
    "B0-55-700H": (2, 2, 0),
    "B0-76-700T": (2, 0, 0),
    # poziomy 3–5
    "B0-08-221Y": (3, 0, 0),
    "B0-76-700U": (3, 0, 0),
    "B0-08-221Z": (4, 0, 0),
    "B0-76-700V": (4, 0, 0),
    "B0-76-700W": (5, 0, 0),
    # hala A: A–E = poziomy 1–5 (pełne palety)
    "A0-01-110A": (1, 0, 0),
    "A0-01-110B": (2, 0, 0),
    "A0-01-110C": (3, 0, 0),
    "A0-01-110D": (4, 0, 0),
    "A3-01-100E": (5, 0, 0),
}


class LetterSlotTests(SimpleTestCase):
    def test_case_table(self):
        for code, want in CASES.items():
            with self.subTest(code=code):
                self.assertEqual(code_slot(code), LetterSlot(*want))

    def test_letter_with_half_suffix_and_lowercase(self):
        self.assertEqual(letter_slot("B0", "c-2"), LetterSlot(1, 2, 2))
        self.assertEqual(letter_slot("b0", "d", half=1), LetterSlot(1, 3, 1))

    def test_unknown_letters_are_none(self):
        # Legacy litery generatora (J/K/L/M/N/O), E poza halą A, F–Z poza A–E w hali A.
        for zone, letter in (("B0", "J"), ("B0", "O"), ("B0", "E"), ("A0", "X"), ("A1", "F"),
                             ("B0", ""), ("B0", "AB")):
            with self.subTest(zone=zone, letter=letter):
                self.assertIsNone(letter_slot(zone, letter))
        self.assertEqual(letter_level("B0", "K", default=7), 7)

    def test_generator_four_part_code_is_not_ewm(self):
        self.assertIsNone(code_slot("B0-01-100-2X"))     # litera = kolumna, poziom jawny

    def test_hall_b_is_default_for_other_zones(self):
        self.assertEqual(letter_level("C1", "X"), 2)
        self.assertEqual(letter_level("", "D"), 1)
        self.assertEqual(letter_level("AB", "B"), 1)     # „AB” to nie strefa hali A (A+cyfry)

    def test_letter_order_matches_physical_order(self):
        # Kolejność od podłogi w addressing.LETTER_ORDER zgodna z poziomami z ewm_levels
        # (S/T to alternatywy A/X w innym typie regału — stąd tylko poziomy niemalejące).
        levels = [letter_slot("B0", c).level for c in LETTER_ORDER]
        self.assertEqual(levels, sorted(levels))
        self.assertLess(LETTER_ORDER.index("B"), LETTER_ORDER.index("C"))
        self.assertLess(LETTER_ORDER.index("C"), LETTER_ORDER.index("D"))
        self.assertLess(LETTER_ORDER.index("D"), LETTER_ORDER.index("G"))
        self.assertLess(LETTER_ORDER.index("G"), LETTER_ORDER.index("H"))
        self.assertLess(LETTER_ORDER.index("H"), LETTER_ORDER.index("Y"))

    def test_height_keys_and_backfill(self):
        self.assertEqual(level_height_keys("B0", "C", 1), ["1C", "1"])
        self.assertEqual(level_height_keys("B0", "C-1", 1), ["1C", "1"])
        self.assertEqual(level_height_keys("B0", "G", 2), ["2G", "2"])
        self.assertEqual(level_height_keys("B0", "X", 2), ["2"])
        self.assertEqual(level_height_keys("A0", "B", 2), ["2"])
        lh = {"1": 2400, "1C": 700, "2": 1800, "3": 1600}
        self.assertEqual(level_height_mm(lh, "B0", "C", 1), 700)     # własna wysokość półki
        self.assertEqual(level_height_mm(lh, "B0", "D", 1), 2400)    # brak „1D” → cały poziom 1
        self.assertEqual(level_height_mm(lh, "B0", "Z", 4), 2400)    # brak „4” → max (jak dotąd)
        self.assertEqual(level_height_mm({}, "B0", "A", 1), 0)
        self.assertEqual(level_height_mm(None, "B0", "A", 1), 0)


class ParsersUseSingleMappingTests(SimpleTestCase):
    def test_parse_loc_code_one_column_per_stack(self):
        for letter, level in (("A", 1), ("B", 1), ("C", 1), ("D", 1), ("X", 2), ("G", 2),
                              ("H", 2), ("Y", 3), ("Z", 4), ("S", 1), ("T", 2), ("W", 5)):
            with self.subTest(letter=letter):
                self.assertEqual(_parse_loc_code(f"B0-39-300{letter}"),
                                 ("B0-39", "300", letter, 0, level))
        self.assertEqual(_parse_loc_code("B0-07-300C-1"), ("B0-07", "300", "C", 0, 1))
        self.assertEqual(_parse_loc_code("A3-01-100E"), ("A3-01", "100", "E", 0, 5))
        # legacy litera kolumny i format generatora — bez zmian
        self.assertEqual(_parse_loc_code("B0-01-100K"), ("B0-01", "100", "K", 2, 2))
        self.assertEqual(_parse_loc_code("B0-01-100-2B"), ("B0-01", "100", "B", 1, 2))

    def test_model_upload_level_from_letter(self):
        self.assertEqual(_parse_location_code("B0-01-300D")[4], 1)   # dawniej D = 4
        self.assertEqual(_parse_location_code("B0-01-300X")[4], 2)   # dawniej X = 24
        self.assertEqual(_parse_location_code("A0-01-110E")[4], 5)

    def test_snapshot_parse_letter_beats_unreliable_sap_level(self):
        row = {"poziom miejsca skł.": 603}
        geo = parse_location("B0-07-300C-1", row)
        self.assertEqual((geo["level"], geo["col_idx"], geo["col_code"]), (1, 0, "C"))
        self.assertEqual(parse_location("B0-55-700H", {})["level"], 2)
        # jawny poziom generatora wygrywa, litera = kolumna
        geo = parse_location("B0-01-100-2B", {})
        self.assertEqual((geo["level"], geo["col_idx"]), (2, 1))

    def test_import_locations_hall_a(self):
        rows = {r["code"]: r for r in parse_rows(
            "h\r\nA0-01-110B;9;0070;200;1;500\r\nB0-39-300D;9;0052;60;1;500\r\n")}
        self.assertEqual((rows["A0-01-110B"]["level"], rows["A0-01-110B"]["col_idx"]), (2, 0))
        self.assertEqual((rows["B0-39-300D"]["level"], rows["B0-39-300D"]["col_idx"]), (1, 0))


class PlanRacksStackTests(SimpleTestCase):
    def _rack(self, letters, stack="300"):
        # poziom z kodu (None) i pozycja z parsera — jak w komórkach bez zapisanego poziomu
        cells = []
        for letter in letters:
            _a, _s, _c, col_idx, _lvl = _parse_loc_code(f"B0-39-{stack}{letter}")
            cells.append((f"B0-39-{stack}{letter}", col_idx, 0, None))
        racks, _rep = plan_racks(cells, {}, slot_mm=1000)
        self.assertEqual(len(racks), 1)
        return racks[0]

    def test_bcd_xyz_stack_is_one_column_four_levels(self):
        rack = self._rack("BCDXYZ")
        self.assertEqual((rack["n_bays"], rack["n_levels"]), (1, 4))
        self.assertEqual(rack["width"], 1.0)                  # jedna kolumna, nie 4

    def test_bcgh_yz_stack(self):
        # B0-55-700: B → C → G → H (+ Y, Z) — G/H to X podzielone pionowo na poziomie 2.
        rack = self._rack("BCGHYZ", stack="700")
        self.assertEqual((rack["n_bays"], rack["n_levels"]), (1, 4))
        self.assertEqual(rack["width"], 1.0)

    def test_compact_stack_five_levels(self):
        self.assertEqual(self._rack("STUVW", stack="700")["n_levels"], 5)


def _user():
    u = get_user_model().objects.create_user(username="ewm-lv", password="x")
    for g in ALL_GROUPS:
        u.groups.add(Group.objects.get_or_create(name=g)[0])
    return u


class CombinedUploadLevelsTests(TestCase):
    """Upload kombinowany: poziom z litery, jedna kolumna mapy na stos, wysokość ze słownika."""
    HEADER = "Miejsce składowania,Typ magazynu,Poziom miejsca skł.\n"
    ROWS = [("B0-39-300B", "0010", 1), ("B0-39-300C", "0010", 2), ("B0-39-300D", "0010", 3),
            ("B0-39-300X", "0010", 4), ("B0-39-300Y", "0010", 5), ("B0-39-300Z", "0010", 6),
            ("B0-55-700B", "0010", 1), ("B0-55-700C", "0010", 1), ("B0-55-700G", "0010", 7),
            ("B0-55-700H", "0010", 8), ("B0-55-700Y", "0010", 9), ("B0-55-700Z", "0010", 9),
            ("B0-07-300C-1", "0010", 603), ("B0-07-300C-2", "0010", 604),
            ("A0-01-110E", "0070", 1)]

    @classmethod
    def setUpTestData(cls):
        cls.user = _user()
        cls.layout = WarehouseLayout.objects.create(name="EWM", is_active=True)
        rt = WarehouseRackType.objects.get(code="0010")
        rt.level_heights = {"1": 2400, "1C": 700, "2": 1800, "2G": 900, "3": 1600, "4": 1500}
        rt.save()
        a = WarehouseRackType.objects.get(code="0070")
        a.level_heights = {"1": 1500, "5": 1200}
        a.save()

    def setUp(self):
        self.client.force_login(self.user)
        body = self.HEADER + "".join(f"{c},{t},{lv}\n" for c, t, lv in self.ROWS)
        f = SimpleUploadedFile("ewm.csv", body.encode("utf-8"), content_type="text/csv")
        self.client.post(reverse("ui:warehouse_combined_upload", args=[self.layout.pk]), {"file": f})
        self.m = {r.location_code: r for r in WarehouseLocationMaster.objects.all()}
        self.c = {c.location_code: c for c in WarehouseLayoutCell.objects.filter(layout=self.layout)}

    def test_levels_from_letter_not_column(self):
        want = {"B0-39-300B": 1, "B0-39-300C": 1, "B0-39-300D": 1, "B0-39-300X": 2,
                "B0-39-300Y": 3, "B0-39-300Z": 4, "B0-55-700G": 2, "B0-55-700H": 2,
                "B0-07-300C-1": 1, "B0-07-300C-2": 1, "A0-01-110E": 5}
        for code, lvl in want.items():
            with self.subTest(code=code):
                self.assertEqual(self.m[code].level, lvl)
                self.assertEqual(self.c[code].level, lvl)

    def test_stack_is_one_map_column(self):
        cols = {self.c[f"B0-39-300{x}"].grid_col for x in "BCDXYZ"}
        self.assertEqual(len(cols), 1)                       # B/C/D nie obok siebie
        self.assertEqual({self.c[f"B0-55-700{x}"].grid_col for x in "BCGHYZ"}, {0})
        self.assertEqual(self.c["B0-07-300C-1"].grid_col, self.c["B0-07-300C-2"].grid_col)

    def test_height_backfill_by_level(self):
        self.assertEqual(self.m["B0-39-300B"].height_mm, 2400)   # „1B” brak → poziom 1
        self.assertEqual(self.m["B0-39-300C"].height_mm, 700)    # własna wysokość półki „1C”
        self.assertEqual(self.m["B0-07-300C-1"].height_mm, 700)
        self.assertEqual(self.m["B0-39-300X"].height_mm, 1800)   # poziom 2 (dawniej brało „2” dla X — ok)
        self.assertEqual(self.m["B0-55-700G"].height_mm, 900)    # „2G”
        self.assertEqual(self.m["B0-55-700H"].height_mm, 1800)   # „2H” brak → poziom 2
        self.assertEqual(self.m["B0-39-300Y"].height_mm, 1600)
        self.assertEqual(self.m["B0-39-300Z"].height_mm, 1500)   # poziom 4 (dawniej „4” = Z też)
        self.assertEqual(self.m["A0-01-110E"].height_mm, 1200)   # hala A: E = poziom 5


class MasterUploadLevelTests(TestCase):
    def test_master_level_from_letter_over_column(self):
        import io

        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Miejsce składowania", "Poziom miejsca skł.", "Typ magazynu"])
        for code, lvl in (("B0-07-300D-2", 606), ("B0-55-700H", 9), ("A1-02-100C", 1), ("XYZ", 3)):
            ws.append([code, lvl, "0010"])
        buf = io.BytesIO()
        wb.save(buf)
        self.client.force_login(_user())
        f = SimpleUploadedFile("m.xlsx", buf.getvalue())
        self.client.post(reverse("ui:warehouse_master_upload"), {"file": f})
        got = dict(WarehouseLocationMaster.objects.values_list("location_code", "level"))
        self.assertEqual(got, {"B0-07-300D-2": 1, "B0-55-700H": 2, "A1-02-100C": 3, "XYZ": 3})


TEMPLATE = Path(__file__).resolve().parents[3] / "excel_templates" / "PalViz_lokalizacje_uklad_wzor.xlsx"


def _template_rows():
    import openpyxl
    ws = openpyxl.load_workbook(TEMPLATE, read_only=True, data_only=True)["Lokalizacje"]
    rows = list(ws.iter_rows(values_only=True))
    return {r[0]: r[4] for r in rows[1:] if r and r[0]}


class ImportTemplateTests(TestCase):
    """Wzór importu (excel_templates) — kolumna „Poziom” = poziom z litery; importery go czytają."""

    def test_template_level_column_matches_letter_rule(self):
        for code, level in _template_rows().items():
            with self.subTest(code=code):
                self.assertEqual(code_slot(code).level, level)

    def test_template_imports_via_list_and_combined_upload(self):
        want = _template_rows()
        self.client.force_login(_user())
        with TEMPLATE.open("rb") as fh:
            self.client.post(reverse("ui:warehouse_layout_list_upload"),
                             {"file": SimpleUploadedFile(TEMPLATE.name, fh.read()), "name": "Wzór"})
        layout = WarehouseLayout.objects.get(name="Wzór")
        self.assertEqual(dict(layout.cells.values_list("location_code", "level")), want)
        with TEMPLATE.open("rb") as fh:
            self.client.post(reverse("ui:warehouse_combined_upload", args=[layout.pk]),
                             {"file": SimpleUploadedFile(TEMPLATE.name, fh.read())})
        got = dict(WarehouseLocationMaster.objects.values_list("location_code", "level"))
        self.assertEqual(got, want)
        # stos B0-39-300 (B/C/D + X/Y/Z) = jedna kolumna mapy
        cols = set(layout.cells.filter(location_code__startswith="B0-39-300")
                   .values_list("grid_col", flat=True))
        self.assertEqual(len(cols), 1)
