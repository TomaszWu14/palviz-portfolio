"""Physical warehouse-map (layout) upload: each cell = a location code at its grid
position; a single cell may hold several depth codes (floor positions)."""
import io

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from ui.models import WarehouseLayout, WarehouseSnapshot
from ui.roles import GROUP_MASTER_DATA


def _md_login(client, username="md"):
    """Create a Master-Data user and log them in (warehouse uploads need that role)."""
    u = get_user_model().objects.create_user(username=username, password="x")
    u.groups.add(Group.objects.get_or_create(name=GROUP_MASTER_DATA)[0])
    client.force_login(u)
    return u


def _xlsx(cells):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r, c, v in cells:
        ws.cell(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class LayoutUploadTests(TestCase):
    def setUp(self):
        _md_login(self.client)

    def _upload(self, cells):
        f = SimpleUploadedFile("lok.xlsx", _xlsx(cells),
                               content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        return self.client.post(reverse("ui:warehouse_layout_upload"), {"file": f, "name": "T"})

    def test_rack_and_multicode_floor_cells_imported(self):
        resp = self._upload([
            (3, 5, "B0-01-300Z"),
            (4, 5, "B0-07-480D-2 | B0-07-480D-1"),   # floor cell: two depth slots
        ])
        self.assertEqual(resp.status_code, 302)
        lay = WarehouseLayout.objects.latest("id")
        codes = set(lay.cells.values_list("location_code", flat=True))
        self.assertEqual(codes, {"B0-01-300Z", "B0-07-480D-2", "B0-07-480D-1"})
        # both floor codes share the same physical grid cell
        d1 = lay.cells.get(location_code="B0-07-480D-1")
        d2 = lay.cells.get(location_code="B0-07-480D-2")
        self.assertEqual((d1.grid_row, d1.grid_col), (d2.grid_row, d2.grid_col))

    def test_level_from_suffix_letter(self):
        self._upload([(3, 5, "B0-01-300Z"), (3, 6, "B0-01-300A")])
        lay = WarehouseLayout.objects.latest("id")
        self.assertEqual(lay.cells.get(location_code="B0-01-300Z").level, 4)  # Z → level 4
        self.assertEqual(lay.cells.get(location_code="B0-01-300A").level, 1)  # A → level 1

    def test_editor_embeds_locations_as_array(self):
        """Regression: locations were double-JSON-encoded → JSON.parse returned a string,
        not an array → editor showed 0 locations (black screen)."""
        self._upload([(3, 5, "B0-01-300Z")])
        resp = self.client.get(reverse("ui:warehouse_editor"))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode()
        self.assertIn("B0-01-300Z", body)
        self.assertIn('"location_code"', body)          # real array element key
        self.assertNotIn('\\"location_code\\"', body)    # NOT double-escaped string

    def test_editor_has_undo_controls(self):
        """Poza-top10: edytor 2D ma undo/redo (Ctrl+Z / Ctrl+Y) — nie tylko
        rollback przy błędzie serwera."""
        self._upload([(3, 5, "B0-01-300Z")])
        body = self.client.get(reverse("ui:warehouse_editor")).content.decode()
        self.assertIn('id="undoBtn"', body)
        self.assertIn('id="redoBtn"', body)
        self.assertIn("function undo(", body)
        self.assertIn("Ctrl+Z", body)


class LayoutSeedTests(TestCase):
    """One-click load of the reference layout bundled in the repo (ui/data/warehouse_layout_b0.xlsx)."""

    def setUp(self):
        _md_login(self.client)

    def test_seed_loads_bundled_layout_as_active(self):
        resp = self.client.post(reverse("ui:warehouse_layout_seed"))
        self.assertEqual(resp.status_code, 302)
        lay = WarehouseLayout.objects.get(name="Układ B0 (wzorcowy)")
        self.assertTrue(lay.is_active)
        # The bundled B0 map has thousands of real locations.
        self.assertGreater(lay.location_count, 5000)
        # Levels from the code suffix (ewm_levels): A and the B/C/D shelves = level 1,
        # X = 2, Y = 3, Z = 4 (B/C/D are shelves INSIDE level 1, not separate levels).
        self.assertEqual(set(lay.cells.values_list("level", flat=True)), {1, 2, 3, 4})
        # Stacks collapse onto 15 parallel aisle rows (not a flat carpet of 100+ rows).
        self.assertEqual(lay.cells.values_list("grid_row", flat=True).distinct().count(), 15)
        # The seed also creates an active master batch so the 3D rack editor can render
        # the empty B0 structure (it builds its rack library from master data).
        from ui.models import WarehouseLocationMasterBatch
        batch = WarehouseLocationMasterBatch.objects.get(name="Układ B0 (wzorcowy)")
        self.assertTrue(batch.is_active)
        self.assertEqual(batch.location_count, lay.location_count)

    def test_seed_master_dims_half_width_for_split_shelves(self):
        # C/D shelves are split 40 cm half-slots; A/B/X/Y/Z are full 80 cm pallet slots.
        self.client.post(reverse("ui:warehouse_layout_seed"))
        from ui.models import WarehouseLocationMaster
        half = WarehouseLocationMaster.objects.filter(location_code__regex=r"-[0-9]+[CD](-[0-9]+)?$")
        full = WarehouseLocationMaster.objects.filter(location_code__regex=r"-[0-9]+[ABXYZ]$")
        self.assertTrue(half.exists(), "no C/D split locations found in bundled B0 data")
        self.assertTrue(full.exists())
        self.assertEqual(set(half.values_list("width_mm", flat=True)), {400})
        self.assertEqual(set(full.values_list("width_mm", flat=True)), {800})
        # Volume must follow the real footprint (40 cm slot carries less).
        self.assertLess(half.first().max_volume_m3, full.first().max_volume_m3)

    def test_seed_deactivates_previous_layouts(self):
        old = WarehouseLayout.objects.create(name="stary", is_active=True)
        self.client.post(reverse("ui:warehouse_layout_seed"))
        old.refresh_from_db()
        self.assertFalse(old.is_active)


def _snapshot_xlsx(headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class SnapshotUploadTests(TestCase):
    XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def setUp(self):
        _md_login(self.client)

    def _upload_snapshot(self, headers, rows):
        f = SimpleUploadedFile("s.xlsx", _snapshot_xlsx(headers, rows), content_type=self.XLSX)
        return self.client.post(reverse("ui:warehouse_map_upload"), {"file": f, "name": "S"})

    def test_handling_unit_marks_location_occupied(self):
        """A stock-at-location export (location + handling unit, no empty/blocked column)
        must read each row as OCCUPIED — not silently default everything to empty."""
        resp = self._upload_snapshot(
            ["Miejsce składowania", "Jednostka obsługi"],
            [["B0-01-300Z", "11111717"], ["B0-01-301Z", "11111718"]])
        self.assertEqual(resp.status_code, 302)
        snap = WarehouseSnapshot.objects.latest("id")
        self.assertEqual(snap.occupied_count, 2)

    def test_ewm_x_flags_mark_blocked_locations(self):
        """Eksport miejsc składowania EWM: blokady to „X” (jak „Puste miejsce skład.”)."""
        self._upload_snapshot(
            ["Miejsce składowania", "Blok. wyd. z magaz.", "Blokada um. w magaz.", "Puste miejsce skład."],
            [["B0-01-100A", "X", "", ""], ["B0-01-200A", "", "X", "X"], ["B0-01-300A", "", "", ""]])
        rows = {r.location_code: r for r in WarehouseSnapshot.objects.latest("id").rows.all()}
        self.assertEqual((rows["B0-01-100A"].blocked_pick, rows["B0-01-100A"].blocked_put), (True, False))
        self.assertEqual((rows["B0-01-200A"].blocked_pick, rows["B0-01-200A"].blocked_put), (False, True))
        self.assertEqual((rows["B0-01-300A"].blocked_pick, rows["B0-01-300A"].blocked_put), (False, False))
        self.assertTrue(rows["B0-01-200A"].is_empty)

    def test_detail_embeds_locations_as_array(self):
        """Regression: detail map double-encoded its JSON → black 3D canvas."""
        self._upload_snapshot(["Miejsce składowania", "Jednostka obsługi"],
                              [["B0-01-300Z", "11111717"]])
        snap = WarehouseSnapshot.objects.latest("id")
        resp = self.client.get(reverse("ui:warehouse_map_detail", args=[snap.id]))
        self.assertEqual(resp.status_code, 200)
        # UX #5: mapa 3D nie może cicho paść — spinner + guard WebGL obecne.
        self.assertContains(resp, "wh3d-status")
        self.assertContains(resp, "__wh3dNoWebGL")
        # R2: lokalizacje nie są już inline w HTML — realne tablice sprawdzamy w JSON
        # (double-encoding niemożliwy: JsonResponse serializuje surowe listy raz).
        locs = self.client.get(reverse("ui:warehouse_map_detail", args=[snap.id]),
                               {"fmt": "json"}).json()["locs"]
        rec = next(r for r in locs if r[9] == "B0-01-300Z")
        self.assertIsInstance(rec, list)            # realna tablica, nie string
        # Bez layoutu pozycje wyliczone z kodu → wpis fizyczny (px/pz numeryczne).
        self.assertIsInstance(rec[0], (int, float))
        self.assertNotIsInstance(rec[0], str)


class SnapshotLayoutBindingTests(TestCase):
    """A snapshot must bind to the active layout (showing rack structure) even when its
    stack still carries the column letter / sub-shelf suffix (e.g. '301C' vs layout '301')."""
    XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def setUp(self):
        _md_login(self.client)
        from ui.models import WarehouseLayout, WarehouseLayoutCell
        lay = WarehouseLayout.objects.create(name="L", is_active=True)
        for stack in ("300", "301", "302"):
            for lvl, suf in ((1, "A"), (4, "X")):
                WarehouseLayoutCell.objects.create(
                    layout=lay, location_code=f"B0-01-{stack}{suf}",
                    grid_row=0, grid_col=int(stack) - 300, level=lvl)

    def test_snapshot_binds_despite_stack_suffix(self):
        # Two of three stacks arrive as sub-shelf codes → stored stack '301C'/'302C'.
        # Without digit-normalised matching only 1/3 would match (<0.5) → flat carpet.
        rows = [["B0-01-300A", "1"], ["B0-01-301C-1", "2"], ["B0-01-302C-2", "3"]]
        f = SimpleUploadedFile("s.xlsx", _snapshot_xlsx(["Miejsce składowania", "Jednostka obsługi"], rows),
                               content_type=self.XLSX)
        self.client.post(reverse("ui:warehouse_map_upload"), {"file": f, "name": "S"})
        snap = WarehouseSnapshot.objects.latest("id")
        resp = self.client.get(reverse("ui:warehouse_map_detail", args=[snap.id]))
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode()
        self.assertNotIn("Brak fizycznego układu", body)   # bound to layout, not the carpet


class Editor3DSmokeTests(TestCase):
    def setUp(self):
        _md_login(self.client)

    def test_editor3d_embeds_rack_library_as_array(self):
        from ui.models import WarehouseLayout, WarehouseLayoutCell
        lay = WarehouseLayout.objects.create(name="L", is_active=True)
        WarehouseLayoutCell.objects.create(layout=lay, location_code="B0-01-300A",
                                            grid_row=3, grid_col=5, level=1)
        resp = self.client.get(reverse("ui:warehouse_editor3d"))
        self.assertEqual(resp.status_code, 200)
        # rack_library must be a real JSON array/object, not a double-encoded string
        self.assertNotIn('id_rack_lib">"', resp.content.decode())

    def test_seeded_layout_renders_racks_in_3d_editor(self):
        # Seeding the bundled B0 layout must also make the empty-rack 3D editor non-empty
        # (it builds its rack library from the master batch the seed now creates).
        self.client.post(reverse("ui:warehouse_layout_seed"))
        resp = self.client.get(reverse("ui:warehouse_editor3d"))
        self.assertEqual(resp.status_code, 200)
        racks = resp.context["rack_library_json"]
        self.assertTrue(racks, "3D editor rack library is empty")
        # Aisles with C/D shelves must flag the split so the floor renders 40 cm half-slots.
        self.assertTrue(any(r.get("hasSplit") for r in racks))
