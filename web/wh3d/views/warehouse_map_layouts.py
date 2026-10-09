# Auto-split from the former monolithic views.py. Feature views live in
# sibling modules; shared imports/constants/helpers stay in core.py.
from ui.views.core import (
    _build_rack_cells, _md_role, get_object_or_404, messages, redirect,
    require_POST, WarehouseLayout, WarehouseLayoutCell, WarehouseLocationMaster,
    WarehouseLocationMasterBatch, WarehouseSnapshot,
)
from ui.views.core.ewm_levels import letter_level, letter_slot
import re

@_md_role
@require_POST
def warehouse_layout_upload(request):
    """Parse MAP (2) sheet from Excel and create a WarehouseLayout."""
    f = request.FILES.get("file")
    if not f:
        messages.error(request, "Brak pliku.")
        return redirect("ui:warehouse_map")
    if f.size > 10 * 1024 * 1024:
        messages.error(request, "Plik zbyt duży (max 10 MB).")
        return redirect("ui:warehouse_map")

    name = request.POST.get("name", "").strip() or f.name
    sheet_name = request.POST.get("sheet", "MAP (2)")
    try:
        layout = _build_layout_from_map_xlsx(f, name, sheet_name)
    except Exception as exc:
        messages.error(request, f"Błąd odczytu pliku: {exc}")
        return redirect("ui:warehouse_map")

    if layout is None:
        messages.error(request, "Nie znaleziono lokalizacji pasujących do wzorca (np. B0-01-100A). Sprawdź format pliku.")
        return redirect("ui:warehouse_map")

    messages.success(request, f"Layout wgrany: {layout.location_count} lokalizacji fizycznych.")
    return redirect("ui:warehouse_map")


# Poziom z litery kodu: ui.views.core.ewm_levels (B/C/D = półki poz. 1, X/G/H/T = 2, …).
# Fallback dla liter spoza EWM (legacy J/K = 2, L/M = 3, N/O = 4; reszta = 1).
_MAP_COL_LEVEL = {"J": 2, "K": 2, "L": 3, "M": 3, "N": 4, "O": 4}
# Allow an optional half suffix (…D-1 / …D-2); capture the zone and the level letter.
_MAP_LOC_RE = re.compile(r"^(B\d+)-\d+-\d+([A-Z])(?:-\d+)?$")


def _build_layout_from_map_xlsx(fileobj, name, sheet_name="MAP (2)"):
    """Build a WarehouseLayout from a physical-map xlsx: every cell holding a location
    code becomes a WarehouseLayoutCell at its real spreadsheet position (grid_row = Excel
    row → depth, grid_col = Excel col → width); the level comes from the code suffix.
    Falls back to the first sheet if `sheet_name` is absent. Returns the active layout, or
    None when no codes were found (the empty layout is rolled back)."""
    import openpyxl as _openpyxl
    wb = _openpyxl.load_workbook(fileobj, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        sheet_name = wb.sheetnames[0]
    ws = wb[sheet_name]

    layout = WarehouseLayout.objects.create(name=name)
    bulk, seen = [], set()
    for row in ws.iter_rows(min_row=1):
        for cell in row:
            if not cell.value:
                continue
            # A single physical cell may hold several codes (floor positions in depth),
            # e.g. "B0-07-480D-2 | B0-07-480D-1" — split and import each.
            for token in re.split(r"[|\n;]+", str(cell.value)):
                raw = token.strip("'").strip()
                m = _MAP_LOC_RE.match(raw)
                if m and raw not in seen:
                    seen.add(raw)
                    bulk.append(WarehouseLayoutCell(
                        layout=layout, location_code=raw,
                        grid_row=cell.row, grid_col=cell.column,
                        level=letter_level(m.group(1), m.group(2).upper())
                        or _MAP_COL_LEVEL.get(m.group(2).upper(), 1)))
                    if len(bulk) >= 2000:
                        WarehouseLayoutCell.objects.bulk_create(bulk, ignore_conflicts=True)
                        bulk = []
    if bulk:
        WarehouseLayoutCell.objects.bulk_create(bulk, ignore_conflicts=True)

    layout.location_count = layout.cells.count()
    if layout.location_count == 0:
        layout.delete()
        return None
    # Deactivate previous layouts only after confirming the new one has data.
    WarehouseLayout.objects.exclude(pk=layout.pk).update(is_active=False)
    layout.save()
    return layout


@_md_role
@require_POST
def warehouse_layout_seed(request):
    """One-click load of the bundled reference warehouse layout (B0) → active layout."""
    from pathlib import Path
    # Dane bundlowane zostają w ui/data (W1: przenosiny widoków, nie zasobów).
    from django.apps import apps as _apps
    p = Path(_apps.get_app_config("ui").path) / "data" / "warehouse_layout_b0.xlsx"
    if not p.exists():
        messages.error(request, "Brak wzorcowego pliku układu w repozytorium.")
        return redirect("ui:warehouse_map")
    try:
        with p.open("rb") as fh:
            layout = _build_layout_from_b0_elevation(fh, name="Układ B0 (wzorcowy)")
    except Exception as exc:
        messages.error(request, f"Błąd odczytu wzorcowego układu: {exc}")
        return redirect("ui:warehouse_map")
    if layout is None:
        messages.error(request, "Nie znaleziono lokalizacji w pliku wzorcowym.")
    else:
        messages.success(request, f"Wczytano wzorcowy układ B0: {layout.location_count} lokalizacji (aktywny).")
    return redirect("ui:warehouse_map")


# B0 elevation drawing: each rack stack is one spreadsheet COLUMN and its shelves are
# stacked across rows. The level comes from the suffix via ui.views.core.ewm_levels —
# NOT the row, because the drawing alternates orientation per aisle (some aisles draw the
# floor at the bottom row, the back-to-back neighbour at the top). B/C/D are shelves
# INSIDE level 1 (one above another), X = level 2, Y = 3, Z = 4; -1/-2 halves share
# the footprint. Unknown letters fall back to level 1.
_B0_LOC_RE = re.compile(r"^(B\d+)-(\d+)-(\d+)([A-Z])(?:-\d+)?$")

# C and D shelves are split in half (the -1/-2 sub-slots), so each holds only ~40 cm of
# face width instead of a full ~80 cm europallet position. Map dimensions/capacity per
# suffix so seeded master data (volumes) matches reality and the 3D view scales bays right.
_B0_HALF_SUFFIXES = {"C", "D"}

def _b0_master_dims(suffix):
    """Slot dims + capacity for a B0 location by shelf suffix.

    Returns (width_mm, depth_mm, height_mm, max_volume_m3, max_weight_kg)."""
    half   = suffix in _B0_HALF_SUFFIXES
    width  = 400 if half else 800
    depth  = 1100
    height = 1000 if half else 1800
    volume = round(width * depth * height / 1_000_000_000, 3)   # mm³ → m³
    weight = 500.0 if half else 1000.0
    return width, depth, height, volume, weight


def _build_layout_from_b0_elevation(fileobj, name):
    """Build a WarehouseLayout from the B0 elevation drawing.

    Each stack's six shelves live in one column across consecutive rows; we collapse them
    into a single footprint — grid_col = spreadsheet column (position along the aisle),
    grid_row = aisle index (15 parallel rack rows) — and take the vertical level from the
    code suffix. This renders the real layout (rows of racks, 6 shelves high) instead of a
    flat carpet of one pillar per cell. Returns the active layout, or None when empty."""
    import openpyxl as _openpyxl
    wb = _openpyxl.load_workbook(fileobj, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    parsed, seen = [], set()
    for row in ws.iter_rows():
        for cell in row:
            if not cell.value:
                continue
            for token in re.split(r"[|\n;]+", str(cell.value)):
                raw = token.strip("'").strip().upper()
                m = _B0_LOC_RE.match(raw)
                if not m or raw in seen:
                    continue
                seen.add(raw)
                parsed.append((raw, m.group(2), cell.column,
                               letter_level(m.group(1), m.group(4), default=1), m.group(4)))
    if not parsed:
        return None

    aisle_row = {a: i for i, a in enumerate(sorted({a for _, a, _, _, _ in parsed}))}
    min_col = min(c for _, _, c, _, _ in parsed)

    layout = WarehouseLayout.objects.create(name=name)
    WarehouseLayoutCell.objects.bulk_create([
        WarehouseLayoutCell(layout=layout, location_code=code,
                            grid_row=aisle_row[aisle], grid_col=col - min_col, level=level)
        for code, aisle, col, level, _suf in parsed
    ], batch_size=2000, ignore_conflicts=True)

    layout.location_count = layout.cells.count()
    if layout.location_count == 0:
        layout.delete()
        return None
    WarehouseLayout.objects.exclude(pk=layout.pk).update(is_active=False)
    layout.save()

    # Also seed an active master batch so the 3D rack editor — which builds its rack
    # library from master data — can render the empty B0 structure. Dimensions/volumes
    # are per-location: C/D shelves get the ~40 cm half-slot, the rest a full ~80 cm slot.
    WarehouseLocationMasterBatch.objects.filter(name=name).delete()
    batch = WarehouseLocationMasterBatch.objects.create(name=name)
    masters = []
    for code, _aisle, _col, level, suf in parsed:
        w, d, h, vol, wt = _b0_master_dims(suf)
        masters.append(WarehouseLocationMaster(
            batch=batch, location_code=code, level=level,
            width_mm=w, depth_mm=d, height_mm=h, max_volume_m3=vol, max_weight_kg=wt))
    WarehouseLocationMaster.objects.bulk_create(masters, batch_size=2000, ignore_conflicts=True)
    batch.location_count = batch.locations.count()
    WarehouseLocationMasterBatch.objects.exclude(pk=batch.pk).update(is_active=False)
    batch.is_active = True
    batch.save()

    return layout

@_md_role
@require_POST
def warehouse_layout_delete(request, pk):
    get_object_or_404(WarehouseLayout, pk=pk).delete()
    messages.success(request, "Layout usunięty.")
    return redirect("ui:warehouse_map")

@_md_role
@require_POST
def warehouse_layout_list_upload(request):
    """Upload layout from a location list (Lokalizacja, Przej., Stos, Poziom)."""
    import openpyxl as _openpyxl

    f = request.FILES.get("file")
    if not f:
        messages.error(request, "Brak pliku.")
        return redirect("ui:warehouse_map")
    if f.size > 10 * 1024 * 1024:
        messages.error(request, "Plik zbyt duży (max 10 MB).")
        return redirect("ui:warehouse_map")

    name = request.POST.get("name", "").strip() or f.name

    try:
        wb = _openpyxl.load_workbook(f, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as exc:
        messages.error(request, f"Błąd odczytu pliku: {exc}")
        return redirect("ui:warehouse_map")

    if not rows:
        messages.error(request, "Plik jest pusty.")
        return redirect("ui:warehouse_map")

    # Detect header row — find columns for location, aisle, stack, level
    _LOC_KEYS   = ["lokalizacja", "miejsce", "location", "adres"]
    _AISLE_KEYS = ["przej", "aisle", "aleja"]
    _STACK_KEYS = ["stos", "stack"]
    _LEVEL_KEYS = ["poziom", "level"]

    def _find_col(headers, keys):
        for i, h in enumerate(headers):
            hl = str(h).lower().strip() if h else ""
            if any(k in hl for k in keys):
                return i
        return None

    col_loc   = _find_col(rows[0], _LOC_KEYS)
    col_aisle = _find_col(rows[0], _AISLE_KEYS)
    col_stack = _find_col(rows[0], _STACK_KEYS)
    col_level = _find_col(rows[0], _LEVEL_KEYS)

    if col_loc is None:
        messages.error(request, "Nie znaleziono kolumny z lokalizacją.")
        return redirect("ui:warehouse_map")

    # Parse rows. Litery EWM → ewm_levels (poziom z litery, jedna kolumna na stos);
    # ta mapa to fallback dla liter spoza EWM (legacy kolumny J/K/L/M/N/O).
    COL_CODE_MAP = {
        'X': (0, 2), 'J': (1, 2), 'K': (2, 2),
        'Y': (0, 3), 'L': (1, 3), 'M': (2, 3),
        'Z': (0, 4), 'N': (1, 4), 'O': (2, 4),
    }

    parsed = []  # (location_code, aisle, stack, col_idx, level)
    seen = set()
    for row in rows[1:]:
        if not row or not row[col_loc]:
            continue
        loc = str(row[col_loc]).strip()
        if not loc or loc in seen:
            continue
        seen.add(loc)

        # Parse aisle/stack/level from columns or from location code
        parts = loc.split('-')
        aisle = str(row[col_aisle]).strip() if col_aisle is not None and col_aisle < len(row) and row[col_aisle] else (parts[1] if len(parts) > 1 else '00')
        stack = str(row[col_stack]).strip() if col_stack is not None and col_stack < len(row) and row[col_stack] else (parts[2][:-1] if len(parts) > 2 and parts[2] else '0')
        col_code = parts[2][-1].upper() if len(parts) > 2 and parts[2] else 'A'
        slot = letter_slot(parts[0], col_code)
        if slot:
            # Kod EWM: poziom ZAWSZE z litery (kolumna „Poziom” bywa błędna dla C-1…D-2).
            col_idx, level = 0, slot.level
        else:
            col_idx, lvl_default = COL_CODE_MAP.get(col_code, (0, 1))
            try:
                level = int(row[col_level]) if col_level is not None and col_level < len(row) and row[col_level] else lvl_default
            except (ValueError, TypeError):
                level = lvl_default

        parsed.append((loc, aisle, stack, col_idx, level))

    if not parsed:
        messages.error(request, "Brak danych do wgrania.")
        return redirect("ui:warehouse_map")

    # Assign grid positions:
    # - unique (aisle, stack) → rank → grid_col = rank * 4 + col_idx
    # - unique aisle → rank → grid_row = rank * 3
    from collections import defaultdict
    aisle_ranks = {}
    stack_ranks = {}  # (aisle, stack) → rank within aisle
    aisle_stack_counter = defaultdict(int)

    for loc, aisle, stack, col_idx, level in parsed:
        if aisle not in aisle_ranks:
            aisle_ranks[aisle] = len(aisle_ranks)
        key = (aisle, stack)
        if key not in stack_ranks:
            stack_ranks[key] = aisle_stack_counter[aisle]
            aisle_stack_counter[aisle] += 1

    from django.db import transaction as _tx
    with _tx.atomic():
        layout = WarehouseLayout.objects.create(name=name)

        bulk = []
        for loc, aisle, stack, col_idx, level in parsed:
            aisle_rank = aisle_ranks[aisle]
            stack_rank = stack_ranks[(aisle, stack)]
            grid_col = stack_rank * 4 + col_idx
            grid_row = aisle_rank * 3
            bulk.append(WarehouseLayoutCell(
                layout=layout,
                location_code=loc,
                grid_col=grid_col,
                grid_row=grid_row,
                level=level,
            ))
            if len(bulk) >= 2000:
                WarehouseLayoutCell.objects.bulk_create(bulk, ignore_conflicts=True)
                bulk = []

        if bulk:
            WarehouseLayoutCell.objects.bulk_create(bulk, ignore_conflicts=True)

        layout.location_count = layout.cells.count()
        layout.save()

        # Deactivate others only after new layout is fully written
        WarehouseLayout.objects.exclude(pk=layout.pk).update(is_active=False)

    messages.success(request, f"Layout wgrany z listy: {layout.location_count} lokalizacji.")
    return redirect("ui:warehouse_map")

@_md_role
@require_POST
def warehouse_demo_seed(request):
    """One-click sample warehouse (3 aisles) so 2D/3D can be seen without any data."""
    from django.db import transaction
    columns = ["A", "B", "C"]
    local_col_idx = {c: i for i, c in enumerate(columns)}
    num_levels = 4
    level_heights = [2494, 2500, 2500, 2200]
    try:
        with transaction.atomic():
            layout = WarehouseLayout.objects.create(name="Demo — magazyn przykładowy")
            WarehouseLayout.objects.exclude(pk=layout.pk).update(is_active=False)
            master_batch = WarehouseLocationMasterBatch.objects.create(name="Master — demo")
            WarehouseLocationMasterBatch.objects.exclude(pk=master_batch.pk).update(is_active=False)

            all_cells, all_locs = [], []
            for ai, aisle in enumerate(["01", "02", "03"]):
                stacks = [str(100 + i * 2) for i in range(6)]
                cells, locs = _build_rack_cells(
                    layout, master_batch, zone="B0", aisle=aisle, stacks=stacks,
                    columns=columns, local_col_idx=local_col_idx, num_levels=num_levels,
                    level_heights=level_heights, slot_width=800, slot_depth=1100,
                    wh_type="DEMO", max_weight=1200, max_volume=2.5, base_row=ai * 2,
                )
                all_cells += cells
                all_locs += locs

            WarehouseLayoutCell.objects.bulk_create(all_cells, ignore_conflicts=True)
            WarehouseLocationMaster.objects.bulk_create(all_locs, ignore_conflicts=True)
            layout.location_count = layout.cells.count(); layout.save()
            master_batch.location_count = master_batch.locations.count(); master_batch.save()
    except Exception as exc:
        messages.error(request, f"Nie udało się utworzyć magazynu demo: {exc}")
        return redirect("ui:warehouse_map")

    messages.success(
        request,
        f"Utworzono magazyn demo: {len(all_cells)} lokalizacji "
        f"(3 aleje × 6 stosów × 3 kol. × 4 poz.). Poniżej Edytor 3D.",
    )
    return redirect("ui:warehouse_editor3d")

@_md_role
def warehouse_map_delete(request, pk):
    snapshot = get_object_or_404(WarehouseSnapshot, pk=pk)
    if request.method == "POST":
        snapshot.delete()
        messages.success(request, "Snapshot usunięty.")
    return redirect("ui:warehouse_map")

__all__ = [
    'warehouse_layout_upload',
    'warehouse_layout_seed',
    'warehouse_layout_delete',
    'warehouse_layout_list_upload',
    'warehouse_demo_seed',
    'warehouse_map_delete',
]
