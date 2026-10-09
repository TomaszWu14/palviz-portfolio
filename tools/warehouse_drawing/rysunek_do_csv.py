"""Rysunek hali (PDF wektorowy) + eksport lokalizacji EWM → CSV geometrii regałów
do „Nowy model magazynu” w GROOVE (web/wh3d/model_geometry.py).

    pip install pymupdf openpyxl
    python tools/warehouse_drawing/rysunek_do_csv.py rzut.pdf Lokalizacje_EWM.xlsx -o regaly.csv

Jak czyta rysunek (skalibrowane na arkuszu Kajima A-ED-01-100-00, hala Logistyczna):
  • rama regału = szara (0.46) pozioma kreska ~8,3 pt; rama co gniazdo, rząd = kolumna ram,
  • skala NIE z tabelki (1:500): 0,12371 m/pt z siatki osi (litery co 6,0 m, cyfry co 5,625 m),
  • przejście EWM = pojedynczy rząd; numeracja rośnie od osi 1 (prawa strona arkusza).
Z EWM: liczba gniazd i poziomów per przejście (poziom z litery kodu B0-PP-GGpL).
Rząd, którego rysunek nie domknął (np. przecięty trasą kablową), bierze zasięg od partnera
plecami, jeśli ten ma pełną liczbę gniazd z EWM.
"""
import argparse
import csv
import re
from collections import defaultdict

M_PER_PT = 0.12371
HALL_ORIGIN_PT = (842.0, 916.0)       # lewa i górna ściana hali wysokiego składowania
FRAME_PT, FRAME_COLOR, DEPTH_CM = 8.3, (0.46, 0.46, 0.46), 103
RACK_CLIP_PT = (830, 870, 2140, 2210)
LEVEL_OF = {**dict.fromkeys("ABCDS", 1), **dict.fromkeys("XGHT", 2), **dict.fromkeys("YU", 3),
            **dict.fromkeys("ZV", 4), "W": 5}
CODE = re.compile(r"^B0-(\d\d)-(\d\d)\d([A-Z])(?:-\d)?$")


def _cluster(vals, tol=1.0):
    out = []
    for v in sorted(vals):
        if out and v - out[-1][-1] <= tol:
            out[-1].append(v)
        else:
            out.append([v])
    return [sum(c) / len(c) for c in out]


def rack_lines(pdf_path):
    """[(x_pt, y0_pt, y1_pt, n_bays)] — jedna pozycja na pojedynczy rząd (kawałki scalone)."""
    import pymupdf
    page = pymupdf.open(pdf_path)[0]
    clip = pymupdf.Rect(*RACK_CLIP_PT)
    frames = []
    for dr in page.get_drawings():
        if not dr["rect"].intersects(clip) or tuple(round(c, 2) for c in dr.get("color") or ()) != FRAME_COLOR:
            continue
        for it in dr["items"]:        # rama = wąski prostokąt (czasem linia) szerokości ~8,3 pt
            if it[0] == "re" and abs(it[1].width - FRAME_PT) < 0.3 and it[1].height < 2:
                frames.append((it[1].x0, (it[1].y0 + it[1].y1) / 2))
            elif it[0] == "l" and abs(it[1].y - it[2].y) < 0.2 and abs(abs(it[1].x - it[2].x) - FRAME_PT) < 0.3:
                frames.append((min(it[1].x, it[2].x), it[1].y))
    lines = []
    for x in _cluster([f[0] for f in frames]):
        ys = _cluster([f[1] for f in frames if abs(f[0] - x) <= 1.0])
        bays = sum(1 for a, b in zip(ys, ys[1:], strict=False) if 13.5 < b - a < 31)     # gniazda 1825/2700/3600
        if bays >= 3:
            lines.append((x, ys[0], ys[-1], bays))
    return lines


def ewm_aisles(xlsx_path):
    """{przejście: (liczba gniazd, liczba poziomów)} z eksportu lokalizacji EWM."""
    from openpyxl import load_workbook
    ws = load_workbook(xlsx_path, read_only=True).active
    rows = ws.iter_rows(values_only=True)
    col = next(rows).index("Miejsce składowania")
    bays, levels = defaultdict(set), defaultdict(int)
    for row in rows:
        m = CODE.match(str(row[col] or "").strip())
        if m:
            aisle = int(m.group(1))
            bays[aisle].add(m.group(2))
            levels[aisle] = max(levels[aisle], LEVEL_OF.get(m.group(3), 1))
    return {a: (len(b), levels[a]) for a, b in bays.items()}


def geometry(lines, ewm):
    """Rzędy → wiersze CSV w układzie hali (x w prawo, y w dół, metry od HALL_ORIGIN_PT)."""
    ox, oy = HALL_ORIGIN_PT
    depth = DEPTH_CM / 100
    by_x = sorted(lines, key=lambda ln: ln[0])
    out = []
    for i, (x, y0, y1, drawn) in enumerate(by_x):
        aisle = len(by_x) - i                                    # przejście 1 = najbliżej osi 1
        n_bays, n_levels = ewm.get(aisle, (drawn, 4))
        left, right = (by_x[i - 1] if i else None), (by_x[i + 1] if i + 1 < len(by_x) else None)
        def gap(o, x=x):
            return abs(o[0] - x) * M_PER_PT - depth if o else 99.0
        partner = min((o for o in (left, right) if o and gap(o) < 1.0), key=gap, default=None)
        if partner and drawn < 0.9 * n_bays and partner[3] >= n_bays - 1:
            y0, y1 = partner[1], partner[2]                      # rysunek nie domknął rzędu
        if partner is not None:
            faces_right = partner is left                        # front po stronie korytarza
        else:
            faces_right = gap(right) < gap(left)
        x_left, top, bottom = (x - ox) * M_PER_PT, (y0 - oy) * M_PER_PT, (y1 - oy) * M_PER_PT
        # kąt −90°: szerokość w +y, głębokość w −x (front +x); 90°: szerokość w −y, front −x
        cx, cy, angle = (x_left + depth, top, -90) if faces_right else (x_left, bottom, 90)
        out.append({"zone": "B0", "rack_id": f"{aisle:02d}", "x_m": round(cx, 2), "y_m": round(cy, 2),
                    "angle_deg": angle, "n_bays": n_bays,
                    "bay_width_cm": round((bottom - top) * 100 / n_bays), "depth_cm": DEPTH_CM,
                    "n_levels": n_levels, "rysunek_gniazd": drawn})
    return sorted(out, key=lambda r: r["rack_id"])


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("pdf")
    ap.add_argument("ewm_xlsx")
    ap.add_argument("-o", "--out", default="regaly_z_rysunku.csv")
    a = ap.parse_args()
    rows = geometry(rack_lines(a.pdf), ewm_aisles(a.ewm_xlsx))
    if not rows:
        raise SystemExit("Nie znaleziono ram regałów — inny arkusz? Sprawdź FRAME_* i RACK_CLIP_PT.")
    with open(a.out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, list(rows[0]), delimiter=";")
        w.writeheader()
        w.writerows(rows)
    print(f"{len(rows)} regałów → {a.out}")


if __name__ == "__main__":
    main()
